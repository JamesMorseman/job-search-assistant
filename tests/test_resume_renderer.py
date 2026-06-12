from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from job_search.generation.generator import DocumentGenerator


def sample_resume_json():
    return {
        "professional_summary": "Entry-level civil engineering candidate focused on structural project work.",
        "skills": [
            "RAM Structural System",
            "AutoCAD",
            "Structural Analysis",
            "Steel Design",
            "ASCE 7-22",
            "AISC",
            "Python",
            "SQLite",
        ],
        "education": [
            {
                "degree": "B.S.",
                "major": "Civil Engineering Technology",
                "institution": "Farmingdale State College",
                "location": "Farmingdale, New York",
                "graduation": "Graduated May 2026",
                "honors": ["Dean's List - Spring 2026"],
                "relevant_coursework": [
                    "Structural Design",
                    "Reinforced Concrete Design",
                    "Steel Design",
                    "Soils and Foundations",
                    "Civil Engineering Materials",
                    "Hydraulics",
                    "Surveying",
                ],
            },
            {
                "degree": "A.S.",
                "major": "Engineering Science",
                "institution": "Suffolk County Community College",
                "location": "Selden, New York",
                "graduation": "Graduated January 2025",
                "relevant_coursework": ["Statics"],
            },
        ],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "organization": "Civil Engineering Technology Program",
                "date": "Spring 2026",
                "bullets": ["Built and reviewed RAM structural models."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Software Automation Project",
                "type": "Personal software / automation project",
                "date": "2026",
                "bullets": ["Built Python automation with SQLite and OpenAI integration."],
            }
        ],
        "experience": [
            {
                "employer": "Urban Air Adventure Park",
                "title": "Event Coordination Department Head",
                "dates": "September 2021 - Present",
                "location": "Lake Grove, New York",
                "bullets": ["Coordinated staffing, logistics, training, and event execution."],
            }
        ],
        "activities": [],
        "certifications": ["FE Civil Exam Candidate (August 2026)"],
    }


def make_generator():
    generator = DocumentGenerator(llm_provider=object())
    generator._profile = {
        "identity": {
            "first_name": "James",
            "last_name": "Morseman",
            "email": "Jamesmorseman@gmail.com",
            "phone": "(631) 559-5622",
            "linkedin_url": "https://example.invalid/not-rendered",
            "location": {"city": "Stony Brook", "state": "New York"},
        },
        "education": [
            {
                "institution": "Farmingdale State College",
                "accreditation": {"program": "ABET-accredited"},
            },
            {"institution": "Suffolk County Community College"},
        ],
        "codes_and_standards": {
            "verified_use_or_awareness": [
                "ASCE 7-22 loading criteria",
                "AISC Steel Construction Manual, 15th Edition",
                "ACI 318 awareness",
                "ASTM D854 soil specific gravity testing",
            ]
        },
    }
    return generator


def render_doc(tmp_path):
    path = tmp_path / "resume.docx"
    make_generator().save_docx(sample_resume_json(), str(path))
    return Document(path)


def paragraph_texts(doc):
    return [p.text for p in doc.paragraphs]


def test_header_matches_template_default_contact(tmp_path):
    doc = render_doc(tmp_path)

    assert doc.paragraphs[0].text == "James Morseman"
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.paragraphs[0].runs[0].font.size.pt == 18
    assert doc.paragraphs[1].text == "Jamesmorseman@gmail.com | (631) 559-5622"
    assert "linkedin" not in doc.paragraphs[1].text.lower()
    assert "Stony Brook" not in doc.paragraphs[1].text


def test_education_order_deans_list_and_coursework_section(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)

    education_idx = texts.index("Education")
    farmingdale_idx = next(i for i, text in enumerate(texts) if "Farmingdale State College" in text)
    suffolk_idx = next(i for i, text in enumerate(texts) if "Suffolk County Community College" in text)
    coursework_idx = texts.index("Relevant Coursework")
    engineering_idx = texts.index("Engineering Experience")

    assert education_idx < farmingdale_idx < suffolk_idx < coursework_idx < engineering_idx
    assert any(text.startswith("Farmingdale State College | Farmingdale, New York") for text in texts)
    assert any(text.startswith("Suffolk County Community College | Selden, New York") for text in texts)
    degree_line = next(text for text in texts if "Civil Engineering Technology" in text)
    assert "ABET Accredited" in degree_line
    assert "Dean's List - Spring 2026" in degree_line
    assert not any(text == "Dean's List - Spring 2026" for text in texts)
    assert all(texts.index("Relevant Coursework") > i for i, text in enumerate(texts) if "College" in text)


def test_education_uses_two_line_layout_with_right_aligned_honors(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)

    school_idx = texts.index("Farmingdale State College | Farmingdale, New York\tMay 2026")
    degree_idx = school_idx + 1
    suffolk_idx = texts.index("Suffolk County Community College | Selden, New York\tJanuary 2025")

    assert texts[degree_idx] == "B.S. in Civil Engineering Technology (ABET Accredited)\tDean's List - Spring 2026"
    assert "Suffolk County Community College | Selden, New York\tJanuary 2025" in texts
    assert doc.paragraphs[degree_idx].paragraph_format.left_indent.inches == 0.25
    assert doc.paragraphs[suffolk_idx + 1].paragraph_format.left_indent.inches == 0.25
    assert doc.paragraphs[school_idx].paragraph_format.tab_stops
    assert doc.paragraphs[degree_idx].paragraph_format.tab_stops
    assert round(doc.paragraphs[degree_idx].paragraph_format.tab_stops[0].position.inches, 1) == 7.5


def test_farmingdale_degree_has_single_abet_reference(tmp_path):
    data = sample_resume_json()
    data["education"][0]["degree"] = "B.S. (ABET Accredited)"
    data["education"][0]["honors"] = ["Dean's List - Spring 2026", "ABET Accredited Program"]

    path = tmp_path / "resume.docx"
    make_generator().save_docx(data, str(path))
    degree_line = next(text for text in paragraph_texts(Document(path)) if "Civil Engineering Technology" in text)

    assert degree_line.count("ABET") == 1
    assert "Dean's List - Spring 2026" in degree_line
    assert "ABET Accredited Program" not in degree_line


def test_dates_use_right_aligned_tab_stops(tmp_path):
    doc = render_doc(tmp_path)
    dated = [p for p in doc.paragraphs if "May 2026" in p.text or "Spring 2026" in p.text]

    assert dated
    for paragraph in dated:
        assert "\t" in paragraph.text
        assert paragraph.paragraph_format.tab_stops
        assert round(paragraph.paragraph_format.tab_stops[0].position.inches, 1) == 7.5


def test_rendered_bullets_remove_trailing_pipe_artifacts(tmp_path):
    data = sample_resume_json()
    data["projects"][0]["bullets"] = [
        "Reviewed RAM structural model |",
        "Coordinated markups;",
        "Prepared calculation package:",
        "Checked connection loads,",
    ]

    path = tmp_path / "resume.docx"
    make_generator().save_docx(data, str(path))
    texts = paragraph_texts(Document(path))

    assert "Reviewed RAM structural model" in texts
    assert "Coordinated markups" in texts
    assert "Prepared calculation package" in texts
    assert "Checked connection loads" in texts
    assert not any(text.endswith(("|", ";", ":", ",")) for text in texts if text.startswith(("Reviewed", "Coordinated", "Prepared", "Checked")))


def test_work_experience_preserves_title_employer_location(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)
    work_line = next(text for text in texts if "Event Coordination Department Head" in text)

    assert "Urban Air Adventure Park" in work_line
    assert "September 2021 - Present" in work_line
    assert "Lake Grove, New York" in texts
    assert "Professional Event Operations" not in " ".join(texts)


def test_project_entries_use_template_two_line_structure(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)

    capstone_idx = texts.index("Senior Capstone Project\tSpring 2026")
    jsa_idx = texts.index("Job Search Assistant\t2026")

    assert texts[capstone_idx + 1] == "Structural Design Lead | Civil Engineering Technology Program"
    assert texts[jsa_idx + 1] == "Project Architect / Lead Developer | Personal Project"
    assert doc.paragraphs[capstone_idx].runs[0].bold is True
    assert doc.paragraphs[capstone_idx].runs[-1].italic is True


def test_project_date_fallback_hierarchy(tmp_path):
    data = sample_resume_json()
    data["projects"] = [
        {"name": "Academic Design Project", "role": "Student designer", "bullets": ["Designed project."]},
        {"name": "Job Search Assistant", "role": "Software Automation Project", "type": "Personal software / automation project", "bullets": ["Built automation."]},
        {"name": "Campus Cleanup", "role": "Volunteer", "type": "Volunteer/community project", "bullets": ["Supported cleanup."]},
    ]
    path = tmp_path / "resume.docx"
    make_generator().save_docx(data, str(path))
    texts = paragraph_texts(Document(path))

    assert "Academic Design Project\t2025" in texts
    assert "Job Search Assistant\t2024" in texts
    assert "Campus Cleanup\t2024" in texts


def test_academic_project_organization_and_invalid_date_fallbacks(tmp_path):
    data = sample_resume_json()
    data["projects"] = [
        {
            "name": "Senior Capstone Project",
            "role": "Structural Engineering Lead / Project Leader",
            "date": "Spring 2026",
            "bullets": ["Led structural design."],
        },
        {
            "name": "Bridge Replacement Construction Management Plan",
            "role": "Academic Project",
            "date": "Academic Project",
            "bullets": ["Prepared construction management planning content."],
        },
    ]
    path = tmp_path / "resume.docx"
    make_generator().save_docx(data, str(path))
    texts = paragraph_texts(Document(path))

    capstone_idx = texts.index("Senior Capstone Project\tSpring 2026")
    bridge_idx = texts.index("Bridge Replacement Construction Management Plan\t2025")

    assert texts[capstone_idx + 1] == "Structural Engineering Lead / Project Leader | Farmingdale State College"
    assert texts[bridge_idx + 1] == "Construction Planning Project | Farmingdale State College"
    assert not any("Academic Project\tAcademic Project" in text for text in texts)


def test_generated_resume_with_profile_work_history_renders_work_experience(tmp_path):
    generator = make_generator()
    generator._profile["experience"] = [
        {
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "start_date": "2021-09",
            "end_date": "present",
            "location": "Lake Grove, NY",
            "bullets": [
                {"text": "Coordinate staffing, logistics, training, and event execution."},
                {"text": "Train new employees and support readiness evaluations."},
            ],
        }
    ]
    data = sample_resume_json()
    data["experience"] = []

    rendered = generator._qa_resume_json(data, "Entry-level structural engineer role with available page capacity.")
    path = tmp_path / "resume.docx"
    generator.save_docx(rendered, str(path))
    texts = paragraph_texts(Document(path))

    assert rendered["experience"]
    assert rendered["rendering_qa"]["page_utilization_estimate"] < 1
    assert "Work Experience" in texts
    assert any("Urban Air Adventure Park" in text for text in texts)


def test_skills_are_grouped_not_single_generic_skills_line(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)

    assert "Technical Skills" in texts
    assert any(text.startswith("Software:") for text in texts)
    assert any(text.startswith("Engineering:") for text in texts)
    assert any(text.startswith("Codes/Standards:") for text in texts)
    codes = next(text for text in texts if text.startswith("Codes/Standards:"))
    assert "ASCE 7" in codes
    assert "AISC Steel Construction Manual" in codes
    assert "ACI 318" in codes
    assert "ASTM D854" in codes
    assert codes.count("ASCE 7") == 1
    assert "ASCE 7-22" not in codes
    assert codes.count("AISC") == 1
    assert not any(text.startswith("Programming/Data:") for text in texts)
    assert not any(text.startswith("Skills:") for text in texts)
    assert "?" not in " ".join(texts)
    assert any("; " in text for text in texts if text.startswith("Software:"))


def test_programming_data_folds_into_software_for_civil_structural_resume():
    generator = make_generator()
    data = sample_resume_json()
    data["professional_summary"] = "Entry-level civil engineer focused on structural design with Python automation support."
    data["skills"] = ["Python", "SQLite", "RAM Structural System", "Structural Analysis"]

    groups = generator._group_skills(
        data["skills"],
        separate_programming_data=generator._is_software_data_heavy_resume(data),
        resume_json=data,
    )

    assert "Programming/Data" not in groups
    assert "Python" in groups["Software"]
    assert "SQLite" in groups["Software"]


def test_programming_data_separates_for_software_data_heavy_resume():
    generator = make_generator()
    data = {
        "professional_summary": "Software engineer and data analyst candidate for a workflow automation role.",
        "keyword_notes": {"included": ["software developer role", "data automation role"]},
        "skills": ["Python", "SQLite", "API integration"],
    }

    groups = generator._group_skills(
        data["skills"],
        separate_programming_data=generator._is_software_data_heavy_resume(data),
        resume_json=data,
    )

    assert "Programming/Data" in groups
    assert "Python" in groups["Programming/Data"]


def test_skill_and_professional_development_spacing_is_compact(tmp_path):
    doc = render_doc(tmp_path)

    compact_labels = ("Software:", "Engineering:", "Codes/Standards:", "Professional Development:")
    compact_paragraphs = [p for p in doc.paragraphs if p.text.startswith(compact_labels)]

    assert compact_paragraphs
    assert all(p.paragraph_format.space_after.pt == 0 for p in compact_paragraphs)


def test_coursework_uses_compact_three_column_table(tmp_path):
    doc = render_doc(tmp_path)
    texts = paragraph_texts(doc)

    assert "Relevant Coursework" in texts
    assert doc.tables
    table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "Structural Design" in table_text
    assert "Hydraulics" in table_text
    assert "Surveying" not in table_text
    assert len(doc.tables[0].columns) == 3


def test_coursework_labels_are_shortened_and_arranged_for_compact_layout(tmp_path):
    data = sample_resume_json()
    data["education"][0]["relevant_coursework"] = [
        "Field Practices in Civil Engineering Technology",
        "Reinforced Concrete Design",
        "Soil Mechanics and Foundations",
        "Structural Analysis",
        "Construction Management",
        "Steel Design",
    ]

    path = tmp_path / "resume.docx"
    make_generator().save_docx(data, str(path))
    table_cells = [cell.text for row in Document(path).tables[0].rows for cell in row.cells]

    assert "Field Practices in Civil Engineering" in table_cells
    assert "Field Practices in Civil Engineering Technology" not in table_cells
    assert "Soils & Foundations" in table_cells
    assert len(DocumentGenerator._arrange_coursework_for_table(data["education"][0]["relevant_coursework"])[:6]) == 6
    assert all(len(text) <= 36 for text in table_cells if text)


def test_renderer_qa_checks_expected_structure():
    qa = make_generator().resume_renderer_qa(sample_resume_json())

    assert qa["sections"].index("Education") < qa["sections"].index("Relevant Coursework")
    assert qa["sections"].index("Relevant Coursework") < qa["sections"].index("Engineering Experience")
    assert "page_utilization_estimate" in qa
    assert "Education" in qa["major_sections_present"]

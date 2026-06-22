"""Tests for Phase 1 document-generation audit layer."""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from job_search.generation.audit import (
    DocumentAuditFailure,
    audit_cover_letter,
    audit_generated_documents,
    audit_report_markdown,
    audit_resume,
    cover_docx_formatting_metadata,
)
from job_search.models import CanonicalJob
from job_search.reporting.selection import SelectionProcessor


def profile(**identity_overrides):
    identity = {
        "first_name": "James",
        "last_name": "Morseman",
        "phone": "(631) 559-5622",
        "email": "Jamesmorseman@gmail.com",
        "linkedin_url": "https://www.linkedin.com/in/james-morseman-82a449344/",
        "github_url": "https://github.com/JamesMorseman",
    }
    identity.update(identity_overrides)
    return {"identity": identity}


def job(company: str = "Acme [Infrastructure Group]") -> CanonicalJob:
    return CanonicalJob(
        source="test",
        source_job_id="job1",
        company=company,
        title="Structural Engineer",
        location_city="New York",
        location_state="NY",
        description_normalized="Structural engineering role.",
    )


def valid_resume_json() -> dict:
    return {
        "professional_summary": "Civil engineering candidate with structural project experience.",
        "skills": ["Structural Analysis", "RAM Structural System", "Steel Design", "AutoCAD"],
        "education": [{"institution": "Farmingdale State College", "degree": "B.S.", "major": "Civil Engineering Technology"}],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "date": "Spring 2026",
                "bullets": ["Performed capstone structural modeling and load-path review."],
            }
        ],
        "experience": [],
        "certifications": [],
    }


def valid_cover_json() -> dict:
    return {
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            "I am interested in the Structural Engineer role at Acme [Infrastructure Group].",
            "My senior capstone engineering work included structural modeling and documentation.",
            "I bring leadership, coordination, and automation initiative to engineering teams.",
        ],
        "closing": "Sincerely,\nJames Morseman",
    }


def valid_resume_text() -> str:
    return "\n".join([
        "James Morseman",
        "(631) 559-5622 | Jamesmorseman@gmail.com",
        "https://www.linkedin.com/in/james-morseman-82a449344/ | https://github.com/JamesMorseman",
        "Professional Summary",
        "Education",
        "Engineering Experience",
        "Senior Capstone Project",
        "Technical Skills",
        "Structural Analysis RAM Structural System Steel Design AutoCAD",
    ])


def valid_cover_text() -> str:
    return "\n".join([
        "James Morseman",
        "(631) 559-5622 | Jamesmorseman@gmail.com | Stony Brook, New York",
        "2026-06-12",
        "Acme [Infrastructure Group]",
        "New York, NY",
        "Dear Hiring Manager,",
        *valid_cover_json()["body_paragraphs"],
        "Sincerely,",
        "James Morseman",
    ])


def failure_codes(result) -> set[str]:
    return {failure["code"] for failure in result.to_dict()["failures"]}


def test_resume_audit_passes_for_valid_resume_header_and_sections():
    result = audit_resume(
        resume_json=valid_resume_json(),
        profile=profile(),
        resume_text=valid_resume_text(),
    )

    assert result.status.value == "PASS"
    assert not result.failures


def test_resume_audit_fails_when_github_missing():
    result = audit_resume(
        resume_json=valid_resume_json(),
        profile=profile(github_url=""),
        resume_text=valid_resume_text().replace(" | https://github.com/JamesMorseman", ""),
    )

    assert "MISSING_GITHUB" in failure_codes(result)
    assert "MALFORMED_HEADER" in failure_codes(result)


def test_resume_audit_fails_when_linkedin_missing():
    result = audit_resume(
        resume_json=valid_resume_json(),
        profile=profile(linkedin_url=""),
        resume_text=valid_resume_text().replace("https://www.linkedin.com/in/james-morseman-82a449344/ | ", ""),
    )

    assert "MISSING_LINKEDIN" in failure_codes(result)


def test_resume_audit_fails_when_capstone_evidence_missing():
    data = valid_resume_json()
    data["projects"] = [{"name": "Bridge Project", "bullets": ["Modeled a bridge."]}]

    result = audit_resume(resume_json=data, profile=profile(), resume_text=valid_resume_text().replace("Capstone", "Bridge"))

    assert "MISSING_CAPSTONE_EVIDENCE" in failure_codes(result)


def test_resume_audit_fails_on_unresolved_placeholder_and_missing_required_section():
    data = valid_resume_json()
    data["education"] = []
    data["professional_summary"] = "Interested in {company}."

    result = audit_resume(resume_json=data, profile=profile(), resume_text=valid_resume_text())

    assert "MISSING_EDUCATION" in failure_codes(result)
    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)


def test_resume_audit_fails_on_leftover_boilerplate_text():
    data = valid_resume_json()
    data["professional_summary"] = "Civil engineering candidate. TODO: add more detail."

    result = audit_resume(resume_json=data, profile=profile(), resume_text=valid_resume_text())

    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)


def test_resume_audit_passes_without_leftover_boilerplate_text():
    result = audit_resume(
        resume_json=valid_resume_json(),
        profile=profile(),
        resume_text=valid_resume_text(),
    )

    assert "PLACEHOLDER_LEAKAGE" not in failure_codes(result)


def test_resume_audit_fails_on_duplicate_bullet_text():
    data = valid_resume_json()
    data["projects"] = [
        {
            "name": "Senior Capstone Project",
            "bullets": ["Performed capstone structural modeling and load-path review."],
        },
        {
            "name": "Second Project",
            "bullets": ["Performed capstone structural modeling and load-path review."],
        },
    ]

    result = audit_resume(resume_json=data, profile=profile(), resume_text=valid_resume_text())

    assert "DUPLICATE_BULLET_TEXT" in failure_codes(result)


def test_resume_audit_passes_with_distinct_bullet_text():
    data = valid_resume_json()
    data["projects"] = [
        {
            "name": "Senior Capstone Project",
            "bullets": ["Performed capstone structural modeling and load-path review."],
        },
        {
            "name": "Second Project",
            "bullets": ["Coordinated steel connection design for a separate structure."],
        },
    ]

    result = audit_resume(resume_json=data, profile=profile(), resume_text=valid_resume_text())

    assert "DUPLICATE_BULLET_TEXT" not in failure_codes(result)


def test_resume_audit_fails_when_page_count_exceeded():
    data = valid_resume_json()
    data["rendering_qa"] = {"estimated_page_count": 2}

    result = audit_resume(
        resume_json=data,
        profile=profile(),
        resume_text=valid_resume_text(),
    )

    assert "PAGE_COUNT_EXCEEDED" in failure_codes(result)


def test_cover_letter_audit_passes_for_valid_business_letter():
    result = audit_cover_letter(
        cover_json=valid_cover_json(),
        profile=profile(),
        job=job(),
        cover_text=valid_cover_text(),
    )

    assert result.status.value == "PASS"
    assert not result.failures


def test_cover_letter_audit_warns_when_rendered_formatting_metadata_missing():
    result = audit_cover_letter(
        cover_json=valid_cover_json(),
        profile=profile(),
        job=job(),
        cover_text=valid_cover_text(),
        formatting_metadata={
            "line_spacing_configured": False,
            "body_paragraph_spacing_configured": False,
            "closing_signature_separation_configured": False,
            "distinct_contact_date_body_blocks": False,
        },
    )

    warning_codes = {warning["code"] for warning in result.to_dict()["warnings"]}

    assert result.status.value == "PASS"
    assert "COVER_LINE_SPACING_NOT_CONFIGURED" in warning_codes
    assert "COVER_BODY_SPACING_NOT_CONFIGURED" in warning_codes
    assert "COVER_CLOSING_SIGNATURE_SPACING_NOT_CONFIGURED" in warning_codes
    assert "COVER_BLOCK_STRUCTURE_NOT_DISTINCT" in warning_codes


def test_cover_letter_docx_formatting_metadata_detects_configured_blocks(tmp_path):
    doc = Document()
    for text in valid_cover_text().splitlines():
        if not text:
            continue
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(8)
        if text == "Sincerely,":
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(0)
    path = tmp_path / "cover.docx"
    doc.save(path)

    metadata = cover_docx_formatting_metadata(str(path))
    result = audit_cover_letter(
        cover_json=valid_cover_json(),
        profile=profile(),
        job=job(),
        cover_text=valid_cover_text(),
        formatting_metadata=metadata,
    )

    assert metadata["line_spacing_configured"] is True
    assert metadata["body_paragraph_spacing_configured"] is True
    assert metadata["closing_signature_separation_configured"] is True
    assert metadata["distinct_contact_date_body_blocks"] is True
    assert not [
        warning for warning in result.to_dict()["warnings"]
        if warning["code"].startswith("COVER_") and warning["code"] != "COVER_LOCATION_REVIEW"
    ]


def test_cover_letter_audit_fails_on_closing_bracket_leakage():
    cover = valid_cover_json()
    cover["body_paragraphs"][-1] = "Paragraph three closing. ["

    result = audit_cover_letter(cover_json=cover, profile=profile(), job=job(), cover_text=valid_cover_text() + "\nclosing. [")

    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)
    assert "MALFORMED_CLOSING_FRAGMENT" in failure_codes(result)


def test_cover_letter_audit_fails_on_standalone_closing_schema_label():
    cover = valid_cover_json()

    result = audit_cover_letter(
        cover_json=cover,
        profile=profile(),
        job=job(),
        cover_text=valid_cover_text() + "\nclosing\nSincerely,\nJames Morseman",
    )

    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)
    assert "MALFORMED_CLOSING_FRAGMENT" in failure_codes(result)


def test_cover_letter_audit_fails_on_single_paragraph_collapse():
    cover = valid_cover_json()
    cover["body_paragraphs"] = [" ".join(valid_cover_json()["body_paragraphs"])]

    result = audit_cover_letter(cover_json=cover, profile=profile(), job=job(), cover_text=valid_cover_text())

    assert "SINGLE_PARAGRAPH_COLLAPSE" in failure_codes(result)
    assert "TOO_FEW_BODY_PARAGRAPHS" in failure_codes(result)


def test_cover_letter_audit_fails_when_closing_signature_missing_or_merged():
    cover = valid_cover_json()
    cover["closing"] = "Sincerely, James Morseman"

    result = audit_cover_letter(cover_json=cover, profile=profile(), job=job(), cover_text=valid_cover_text())

    assert "MERGED_CLOSING_SIGNATURE" in failure_codes(result)


def test_cover_letter_audit_allows_legitimate_bracketed_company_name():
    result = audit_cover_letter(
        cover_json=valid_cover_json(),
        profile=profile(),
        job=job(),
        cover_text=valid_cover_text(),
    )

    assert result.status.value == "PASS"


def test_cover_letter_audit_fails_unresolved_template_markers():
    cover = valid_cover_json()
    cover["body_paragraphs"][0] = "I am interested in {company}."

    result = audit_cover_letter(cover_json=cover, profile=profile(), job=job(), cover_text=valid_cover_text())

    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)


def test_cover_letter_audit_fails_on_leftover_boilerplate_text():
    cover = valid_cover_json()
    cover["body_paragraphs"][0] = "Lorem ipsum dolor sit amet, interested in the role."

    result = audit_cover_letter(cover_json=cover, profile=profile(), job=job(), cover_text=valid_cover_text())

    assert "PLACEHOLDER_LEAKAGE" in failure_codes(result)


def test_document_audit_report_serializes_pass_and_fail_outputs():
    passed = audit_generated_documents(
        resume_json=valid_resume_json(),
        cover_letter_json=valid_cover_json(),
        profile=profile(),
        job=job(),
        resume_text=valid_resume_text(),
        cover_letter_text=valid_cover_text(),
    )
    failed = audit_generated_documents(
        resume_json=valid_resume_json(),
        cover_letter_json=valid_cover_json(),
        profile=profile(github_url=""),
        job=job(),
        resume_text=valid_resume_text(),
        cover_letter_text=valid_cover_text() + "\nclosing. [",
    )

    assert passed["status"] == "PASS"
    assert failed["status"] == "FAIL"
    assert any(item["code"] == "MISSING_GITHUB" for item in failed["failures"])
    assert "Document Audit: FAIL" in audit_report_markdown(failed)


def test_upload_gate_blocks_failed_audit_before_upload(tmp_path):
    class FakeGenerator:
        _profile = profile(github_url="")

        def save_docx(self, resume_json, output_path):
            doc = Document()
            for line in valid_resume_text().replace(" | https://github.com/JamesMorseman", "").splitlines():
                doc.add_paragraph(line)
            doc.save(output_path)

        def save_cover_docx(self, cover_json, output_path, job=None, today=None):
            doc = Document()
            for line in valid_cover_text().splitlines():
                doc.add_paragraph(line)
            doc.save(output_path)

    class FakeSheets:
        uploaded = False

        def upload_document(self, local_path, filename, folder_id=None):
            self.uploaded = True
            return f"https://example.invalid/{filename}"

    proc = SelectionProcessor()
    proc._generator = FakeGenerator()
    proc.sheets = FakeSheets()
    row = {
        "canonical_job_id": "job1",
        "source": "test",
        "source_job_id": "job1",
        "firm_id": None,
        "company": "Acme [Infrastructure Group]",
        "title": "Structural Engineer",
        "location_city": "New York",
        "location_state": "NY",
        "description_normalized": "Structural engineering role.",
        "description_raw": "",
        "apply_url": "",
        "ats_type": "unknown",
    }

    with pytest.raises(DocumentAuditFailure) as exc:
        proc._upload_docs(
            row,
            {"resume_json": valid_resume_json(), "cover_letter_json": valid_cover_json()},
            today="2026-06-12",
        )

    assert "MISSING_GITHUB" in str(exc.value)
    assert proc.sheets.uploaded is False

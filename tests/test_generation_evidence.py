"""Tests for evidence packet integration in document generation."""

import json
from pathlib import Path

from docx import Document

from job_search.generation.generator import DocumentGenerator
from job_search.llm.types import LLMResponse
from job_search.models import CanonicalJob

from tests.test_evidence_selection import sample_profile


class FakeLLMProvider:
    provider_name = "fake"

    def __init__(self):
        self.requests = []

    def generate_json(self, request):
        self.requests.append(request)
        if len(self.requests) == 1:
            return LLMResponse(content=json.dumps({
                "professional_summary": "Structural candidate.",
                "skills": ["RAM"],
                "education": [],
                "experience": [],
                "projects": [],
                "activities": [],
                "certifications": [],
                "keyword_notes": {"included": [], "rationale": ""},
            }))
        return LLMResponse(content=json.dumps({
            "salutation": "Dear Hiring Manager,",
            "body_paragraphs": [
                "I am interested in this structural role.",
                "My capstone, leadership, and automation experience align with the role.",
                "I would welcome the opportunity to contribute.",
            ],
            "closing": "Sincerely, James Morseman",
        }))

    def submit_json_batch(self, requests):
        raise NotImplementedError

    def retrieve_batch_status(self, batch_id):
        raise NotImplementedError

    def fetch_json_batch_results(self, batch_id):
        raise NotImplementedError


def make_job(description: str | None = None) -> CanonicalJob:
    return CanonicalJob(
        source="test",
        source_job_id="job1",
        company="Acme",
        title="Structural Engineer",
        description_normalized=description or "Structural design using RAM, steel, concrete, and AutoCAD.",
    )


def user_prompt(fake: FakeLLMProvider, index: int = 0) -> str:
    return fake.requests[index].messages[1].content


def test_generator_uses_selected_evidence_instead_of_full_profile():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    profile = sample_profile()
    profile["unrelated_private_note"] = "DO NOT SEND THIS IRRELEVANT FULL PROFILE NOTE"
    generator._profile = profile

    result = generator.generate(make_job())

    assert result["evidence_fallback_used"] is False
    prompt = user_prompt(fake)
    assert "selected_evidence_packet" in prompt
    assert "baseline_profile_facts" in prompt
    assert "rb_structural" in prompt
    assert "DO NOT SEND THIS IRRELEVANT FULL PROFILE NOTE" not in prompt


def test_generator_includes_work_history_in_baseline_profile_facts():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    profile = sample_profile()
    profile["experience"] = [{
        "employer": "Urban Air Adventure Park",
        "title": "Event Coordination Department Head",
        "start_date": "2021-09",
        "end_date": "present",
        "location": "Lake Grove, NY",
    }]
    generator._profile = profile

    generator.generate(make_job())

    prompt = user_prompt(fake)
    assert "baseline_profile_facts" in prompt
    assert "Urban Air Adventure Park" in prompt
    assert "Event Coordination Department Head" in prompt


def test_resume_style_guide_is_included_in_prompt():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    generator._profile = sample_profile()

    generator.generate(make_job())

    prompt = user_prompt(fake)
    assert "Resume Rendering Spec" in prompt
    assert "Resume Content Rules" in prompt
    assert "Target length: 1 page" in prompt
    assert "Relevant Coursework is conditional" in prompt


def test_cover_letter_style_guide_is_included_in_prompt():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    generator._profile = sample_profile()

    generator.generate(make_job())

    prompt = user_prompt(fake, index=1)
    assert "Cover Letter Style Guide" in prompt
    assert "1 page maximum" in prompt
    assert "Paragraph 1: name the role/company" in prompt
    assert "Do not repeat resume bullets mechanically" in prompt
    assert "Use exactly 3 body paragraphs" in prompt
    assert "capstone/technical engineering evidence" in prompt
    assert "leadership/management plus Job Search Assistant evidence" in prompt
    assert "Do not mention ChatGPT, Codex" in prompt


def test_cover_letter_pipeline_preserves_generated_body_paragraphs():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    generator._profile = sample_profile()

    result = generator.generate(make_job())
    cover = result["cover_letter_json"]

    assert len(fake.requests) == 2
    assert result["document_audit"]["status"] in {"PASS", "FAIL"}
    assert "resume" in result["document_audit"]
    assert "cover_letter" in result["document_audit"]
    assert len(cover["body_paragraphs"]) == 3
    assert cover["rendering_diagnostics"]["raw_body_paragraph_count"] == 3
    assert cover["rendering_diagnostics"]["body_paragraph_count"] == 3
    assert cover["rendering_diagnostics"]["multi_paragraph_output"] is True
    assert cover["rendering_diagnostics"]["closing_signature_separate"] is True


def test_cover_letter_text_rendering_keeps_paragraphs_and_signature_separate():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["Paragraph one.", "Paragraph two.", "Paragraph three."],
        "closing": "Sincerely, James Morseman",
    })

    rendered = generator._render_cover_text(cover)

    assert "Paragraph one.\n\nParagraph two." in rendered
    assert "Paragraph three.\n\nSincerely,\nJames Morseman" in rendered
    assert "body_paragraphs_placeholder" not in rendered


def test_cover_letter_signature_uses_first_last_profile_name():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = {
        "identity": {
            "first_name": "James",
            "middle_name": "Robert",
            "last_name": "Morseman",
        }
    }

    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["Paragraph one.", "Paragraph two.", "Paragraph three."],
        "closing": "Sincerely, James Robert Morseman",
    })

    assert cover["closing"] == "Sincerely,\nJames Morseman"
    assert "James Robert Morseman" not in generator._render_cover_text(cover)


def test_cover_letter_prompt_requires_three_paragraph_content_structure():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    generator._profile = sample_profile()

    generator.generate(make_job(
        "Structural role at Acme requiring RAM, steel design, leadership, data workflows, and reporting."
    ))
    prompt = user_prompt(fake, index=1)

    assert "1. role/company fit" in prompt
    assert "2. capstone/technical engineering evidence" in prompt
    assert "3. leadership/management plus Job Search Assistant evidence when relevant" in prompt
    assert "Do not mention ChatGPT, Codex" in prompt


def test_cover_letter_docx_rendering_uses_multiple_paragraphs_without_placeholders(tmp_path):
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = sample_profile()
    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["Paragraph one.", "Paragraph two.", "Paragraph three."],
        "closing": "Sincerely, James Morseman",
    })
    path = tmp_path / "cover.docx"

    generator.save_cover_docx(cover, str(path), job=make_job(), today="2026-06-12")
    texts = [p.text for p in Document(path).paragraphs if p.text]

    assert "Paragraph one." in texts
    assert "Paragraph two." in texts
    assert "Paragraph three." in texts
    assert "Sincerely," in texts
    assert "James Morseman" in texts
    assert texts.index("Sincerely,") < texts.index("James Morseman")
    assert not any("placeholder" in text.lower() for text in texts)

    doc = Document(path)
    paragraphs = {p.text: p for p in doc.paragraphs if p.text}
    assert paragraphs["Paragraph one."].paragraph_format.line_spacing == 1.15
    assert paragraphs["Paragraph one."].paragraph_format.space_after.pt == 11
    assert paragraphs["Sincerely,"].paragraph_format.space_before.pt == 14
    assert paragraphs["Sincerely,"].paragraph_format.space_after.pt == 0
    assert paragraphs["James Morseman"].paragraph_format.line_spacing == 1.15
    assert paragraphs["James Morseman"].paragraph_format.space_before.pt == 4
    assert paragraphs["James Morseman"].paragraph_format.space_after.pt == 8


def test_cover_letter_renderer_rejects_template_placeholders():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = {
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["body_paragraphs_placeholder"],
        "closing": "Sincerely,\nJames Morseman",
    }

    try:
        generator._render_cover_text(cover)
    except ValueError as exc:
        assert "placeholder or marker leaked" in str(exc).lower()
    else:
        raise AssertionError("Expected placeholder leakage to fail rendering.")


def test_cover_letter_renderer_rejects_unresolved_bracket_markers():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = {
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            "Paragraph one.",
            "Paragraph two.",
            "Paragraph three closing. [",
        ],
        "closing": "Sincerely,\nJames Morseman",
    }

    try:
        generator._render_cover_text(cover)
    except ValueError as exc:
        assert "placeholder or marker leaked" in str(exc).lower()
    else:
        raise AssertionError("Expected unresolved bracket marker to fail rendering.")


def test_cover_letter_renderer_rejects_standalone_closing_schema_label():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = {
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            "Paragraph one.",
            "Paragraph two.",
            "Paragraph three.",
        ],
        "closing": "closing",
    }

    try:
        generator._render_cover_text(cover)
    except ValueError as exc:
        assert "cover_letter_closing" in str(exc)
    else:
        raise AssertionError("Expected standalone closing schema label to fail rendering.")


def test_cover_letter_renderer_allows_legitimate_bracketed_company_name():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            "I am interested in the engineering role at Acme [Infrastructure Group].",
            "Paragraph two connects technical project work to the role.",
            "Paragraph three closes with interest in contributing to the team.",
        ],
        "closing": "Sincerely,\nJames Morseman",
    })

    rendered = generator._render_cover_text(cover)

    assert "Acme [Infrastructure Group]" in rendered


def test_cover_letter_qa_rejects_too_few_body_paragraphs():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = {
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["Paragraph one.", "Paragraph two."],
        "closing": "Sincerely,\nJames Morseman",
    }

    try:
        generator._qa_cover_letter_json(cover)
    except ValueError as exc:
        assert "at least 3 separate body paragraphs" in str(exc)
    else:
        raise AssertionError("Expected short cover letter body to fail QA.")


def test_cover_letter_qa_repairs_closing_merged_into_body():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            "Paragraph one.",
            "Paragraph two.",
            "Paragraph three. Sincerely, James Morseman",
        ],
        "closing": "Sincerely, James Morseman",
    })

    assert cover["body_paragraphs"] == ["Paragraph one.", "Paragraph two.", "Paragraph three."]
    assert cover["closing"] == "Sincerely,\nJames Morseman"


def test_cover_letter_docx_header_has_stable_contact_block(tmp_path):
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = {
        "identity": {
            "first_name": "James",
            "last_name": "Morseman",
            "phone": "(631) 559-5622",
            "email": "Jamesmorseman@gmail.com",
            "linkedin_url": "https://www.linkedin.com/in/james-morseman-82a449344/",
            "location": {"city": "Stony Brook", "state": "New York"},
        }
    }
    cover = generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": ["Paragraph one.", "Paragraph two.", "Paragraph three."],
        "closing": "Sincerely, James Morseman",
    })
    path = tmp_path / "cover.docx"

    generator.save_cover_docx(
        cover,
        str(path),
        include_location=False,
        include_linkedin=True,
    )
    texts = [p.text for p in Document(path).paragraphs if p.text]

    assert texts[0] == "James Morseman"
    assert texts[1] == (
        "(631) 559-5622 | Jamesmorseman@gmail.com | "
        "https://www.linkedin.com/in/james-morseman-82a449344/"
    )
    assert "Stony Brook" not in texts[1]

    doc = Document(path)
    assert doc.paragraphs[0].paragraph_format.line_spacing == 1.15
    assert doc.paragraphs[1].paragraph_format.line_spacing == 1.15
    assert doc.paragraphs[1].paragraph_format.space_after.pt == 16


def test_cover_letter_prompt_includes_leadership_capstone_and_jsa_evidence():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    profile = sample_profile()
    profile["cover_letter_fragment_bank"].extend([
        {
            "id": "cf_leadership",
            "text": "Led event teams through staffing, training, scheduling, and customer communication.",
            "tags": ["leadership", "coordination"],
        },
        {
            "id": "cf_automation",
            "text": "Built the Job Search Assistant to organize job data, document generation, and follow-up workflows.",
            "tags": ["automation", "python"],
        },
    ])
    generator._profile = profile

    generator.generate(make_job(
        "Structural role using RAM and steel design with project coordination, leadership, "
        "Python automation, data workflows, and reporting."
    ))
    prompt = user_prompt(fake, index=1)

    assert "cf_leadership" in prompt
    assert "capstone" in prompt
    assert "personal_jsa" in prompt
    assert "Job Search Assistant" in prompt


def test_resume_qa_expands_underfilled_resume_from_verified_profile_facts():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = {
        "education": [{
            "institution": "Farmingdale State College",
            "relevant_coursework": [
                "Structural Design",
                "Reinforced Concrete Design",
                "Soils, Foundations, and Earth Structures",
                "Civil Engineering Materials",
                "Hydraulics",
                "Transportation Engineering",
                "Field Practices in Civil Engineering Technology",
                "Technical Writing",
                "Engineering Economics",
            ],
        }],
        "resume_bullet_bank": {
            "structural": [
                "Built and analyzed RAM Structural System models for a proposed three-story classroom addition, supporting gravity/lateral load evaluation and preliminary steel framing decisions.",
                "Applied ASCE 7-22 loading criteria and AISC steel design methodology while evaluating framing layouts, tributary areas, line loads, axial loads, and serviceability behavior.",
                "Designed steel beam and column members for a five-story residential building project using ASCE 7-16 LRFD load combinations and AISC Steel Construction Manual checks.",
            ],
            "construction": [
                "Prepared construction management planning content for a bridge replacement project, including schedule logic, traffic control, safety, equipment needs, permitting awareness, and cost discussion.",
            ],
            "leadership": [
                "Led structural scope and coordinated multidisciplinary capstone deliverables, including task delegation, milestone tracking, report integration, and final presentation support.",
            ],
            "software_data_automation": [
                "Built a Python job-search automation platform with multi-source ingestion, SQLite persistence, fit grading, document generation, and a CLI workflow.",
                "Integrated Google Drive and Sheets workflows and maintained an automated test suite with 100+ passing tests.",
            ],
        },
        "technical_skills": {
            "structural": ["Structural analysis and design", "Steel beam design", "Gravity load development"],
            "construction": ["Construction planning", "Project scheduling"],
        },
        "software_tools": {"verified": ["RAM Structural System", "AutoCAD", "Microsoft Excel"]},
        "projects": [{
            "name": "Construction Management Project of Bridge Replacement - Spencer",
            "type": "CLASS_PROJECT",
            "date": "2025",
            "role": "Student project team member",
            "bullets": [
                {
                    "text": (
                        "Reviewed bridge replacement scope, traffic control, schedule logic, "
                        "equipment, cost, safety, and construction-phase considerations."
                    )
                }
            ],
        }],
        "certifications": {
            "professional_development": [
                "Expected next step after FE passage is Engineer-in-Training certification.",
                "Long-term professional goal is Professional Engineer licensure after qualifying experience.",
            ],
        },
        "profile_summary": {
            "current_status": "Recent Civil Engineering Technology graduate preparing for the FE Civil examination.",
            "positioning": (
                "Applied civil engineering graduate with strong structural interest, "
                "capstone leadership, hands-on laboratory exposure, and operations experience."
            ),
            "primary_differentiators": [
                "Structural Engineering Lead and functional project leader for multidisciplinary capstone project.",
                "Built and worked with RAM Structural System models for an educational facility addition project.",
            ],
        },
    }
    resume = {
        "professional_summary": "Civil engineering graduate focused on structural work.",
        "skills": ["RAM Structural System"],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
            "major": "Civil Engineering Technology",
            "graduation": "May 2026",
            "relevant_coursework": ["Structural Design"],
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "dates": "September 2021 - Present",
            "location": "Lake Grove, New York",
            "bullets": ["Coordinated staffing and logistics for event operations."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "date": "Spring 2026",
                "bullets": ["Modeled a proposed Baldwin High School addition."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Automation Project",
                "date": "2026",
                "bullets": ["Built a Python workflow for job-search tracking."],
            },
        ],
        "certifications": ["FE Civil Exam Candidate"],
    }

    original_word_count = generator._resume_word_count(resume)
    cleaned = generator._qa_resume_json(resume, "structural civil construction engineering role")
    qa = generator.resume_renderer_qa(cleaned)

    assert generator._resume_word_count(cleaned) > original_word_count
    assert qa["page_utilization_estimate"] > generator._resume_page_utilization(resume)
    assert cleaned["rendering_allocation"]["one_page_enforced"] is True
    assert "skills_items" in cleaned["rendering_allocation"]["expansion_order"]
    assert "professional_development_items" in cleaned["rendering_allocation"]["expansion_order"]
    assert "summary_detail" in cleaned["rendering_allocation"]["expansion_order"]
    project_names = [project["name"] for project in cleaned["projects"]]
    assert "Bridge Replacement Construction Management Planning" in project_names
    assert "Structural and Construction Coursework" not in project_names


def test_resume_qa_enforces_one_page_policy_with_protected_evidence():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = sample_profile()
    resume = {
        "professional_summary": " ".join(["Civil structural candidate"] * 35),
        "skills": [f"Skill {index}" for index in range(18)],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
            "major": "Civil Engineering Technology",
            "graduation": "May 2026",
            "relevant_coursework": [
                "Structural Design",
                "Reinforced Concrete Design",
                "Soils, Foundations, and Earth Structures",
                "Civil Engineering Materials",
                "Hydraulics",
                "Transportation Engineering",
                "Field Practices in Civil Engineering Technology",
                "Technical Writing",
                "Engineering Economics",
            ],
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "dates": "September 2021 - Present",
            "location": "Lake Grove, New York",
            "bullets": ["Work bullet one.", "Work bullet two.", "Work bullet three."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "date": "Spring 2026",
                "bullets": ["Capstone one.", "Capstone two.", "Capstone three.", "Capstone four.", "Capstone five."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Automation Project",
                "date": "2026",
                "bullets": ["JSA one.", "JSA two.", "JSA three."],
            },
            {
                "name": "Bridge Replacement Project",
                "role": "Student Designer",
                "date": "2025",
                "bullets": ["Bridge one.", "Bridge two.", "Bridge three."],
            },
        ],
        "certifications": [
            "FE Civil Exam Candidate",
            "Expected next step after FE passage is Engineer-in-Training certification.",
            "Long-term professional goal is Professional Engineer licensure after qualifying experience.",
        ],
    }

    cleaned = generator._qa_resume_json(resume, "structural civil engineering role")
    qa = generator.resume_renderer_qa(cleaned)

    assert qa["estimated_page_count"] == 1
    assert "page_count_exceeded" not in qa["warnings"]
    assert cleaned["rendering_allocation"]["one_page_enforced"] is True
    assert any(generator._is_capstone_project(project) for project in cleaned["projects"])
    assert any(generator._is_job_search_assistant_project(project) for project in cleaned["projects"])
    assert cleaned["experience"]
    assert len(generator._selected_coursework(cleaned)) <= 6
    assert len(cleaned["certifications"]) <= 1
    assert len(cleaned["experience"][0]["bullets"]) <= 2


def test_resume_trimming_preserves_accepted_project_identities():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = sample_profile()
    resume = {
        "professional_summary": "Civil engineering candidate with structural, automation, and construction management breadth.",
        "skills": ["RAM Structural System", "Python automation", "SQLite", "Construction planning"],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
            "major": "Civil Engineering Technology",
            "graduation": "May 2026",
            "relevant_coursework": ["Structural Design", "Field Practices in Civil Engineering Technology"],
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Coordinated staffing.", "Improved tracking."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "bullets": ["Capstone one.", "Capstone two.", "Capstone three.", "Capstone four."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Automation Project",
                "bullets": ["JSA one.", "JSA two."],
            },
            {
                "name": "Bridge Replacement Construction Management Planning",
                "role": "Student project team member",
                "bullets": ["Bridge planning one.", "Bridge planning two.", "Bridge planning three."],
            },
        ],
        "certifications": ["FE Civil Exam Candidate"],
    }

    cleaned = generator._qa_resume_json(resume, "entry-level infrastructure role")
    project_names = [project["name"] for project in cleaned["projects"]]

    assert "Senior Capstone Project" in project_names
    assert "Job Search Assistant" in project_names
    assert "Bridge Replacement Construction Management Planning" in project_names
    assert "Structural and Construction Coursework" not in project_names
    assert cleaned["projects"][2]["name"] == "Bridge Replacement Construction Management Planning"
    assert len(cleaned["projects"][2]["bullets"]) <= 1


def test_resume_qa_restores_drifted_construction_management_project_identity():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    resume = {
        "professional_summary": "Civil engineering candidate with construction planning evidence.",
        "skills": ["Construction management planning"],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
        }],
        "experience": [],
        "projects": [{
            "name": "Bridge Replacement Construction Management Plan",
            "role": "Construction Planning Project",
            "bullets": [
                "Prepared construction management planning content for a bridge replacement project, "
                "including schedule logic, traffic control, safety, equipment needs, and cost discussion."
            ],
        }],
        "certifications": [],
    }

    cleaned = generator._qa_resume_json(resume, "construction management planning role")

    assert cleaned["projects"][0]["name"] == "Bridge Replacement Construction Management Planning"


def test_resume_qa_restores_drifted_stormwater_project_identity():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    resume = {
        "professional_summary": "Civil engineering candidate with stormwater evidence.",
        "skills": ["Stormwater runoff calculations"],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
        }],
        "experience": [],
        "projects": [{
            "name": "Campus Stormwater Cistern Design",
            "role": "Academic Hydrology / Stormwater Project",
            "bullets": [
                "Calculated roof runoff and evaluated cistern storage, filtration, pumping, overflow, and irrigation reuse."
            ],
        }],
        "certifications": [],
    }

    cleaned = generator._qa_resume_json(resume, "land development stormwater drainage role")

    assert cleaned["projects"][0]["name"] == "Stormwater Detention / Cistern Design Project"


def test_resume_qa_restores_role_relevant_profile_project_before_supporting_content():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = {
        "projects": [
            {
                "name": "Stormwater Detention / Cistern Design Project",
                "type": "CLASS_PROJECT",
                "date": "2026-05",
                "role": "Student project team member",
                "bullets": [
                    {"text": "Evaluated roof catchment, runoff volume, cistern storage, and overflow routing."},
                    {"text": "Compared detention pond, infiltration trench, green roof, and cistern system alternatives."},
                ],
            },
        ],
        "technical_skills": {"water_site_civil": ["Stormwater design", "Drainage planning"]},
    }
    resume = {
        "professional_summary": "Civil engineering candidate focused on land development and site civil work.",
        "skills": ["Stormwater design"],
        "education": [{"institution": "Farmingdale State College", "degree": "Bachelor of Science"}],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Coordinated staffing.", "Improved tracking."],
        }],
        "projects": [
            {"name": "Senior Capstone Project", "role": "Structural Design Lead", "bullets": ["Capstone one."]},
            {"name": "Job Search Assistant", "role": "Automation Project", "bullets": ["JSA one."]},
        ],
        "certifications": [],
    }

    cleaned = generator._qa_resume_json(
        resume,
        "Civil Engineer land development site civil stormwater drainage role.",
    )
    project_names = [project["name"] for project in cleaned["projects"]]
    order = cleaned["rendering_allocation"]["expansion_order"]
    stormwater_bullet = cleaned["projects"][2]["bullets"][0]

    assert "Stormwater Detention / Cistern Design Project" in project_names
    assert len(cleaned["projects"][2]["bullets"]) == 1
    assert "runoff volume" in stormwater_bullet
    assert "detention pond" in stormwater_bullet
    assert "academic_project_bullets" in order
    if "work_experience_bullets" in order:
        assert order.index("academic_project_bullets") < order.index("work_experience_bullets")
    assert "expansion_order_risk" not in generator.resume_renderer_qa(cleaned)["warnings"]


def test_resume_summary_trim_avoids_sentence_fragments():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    summary = (
        "Recent Civil Engineering Technology graduate preparing for the FE Civil exam. "
        + " ".join(["Applied civil engineering candidate"] * 25)
    )
    data = {
        "professional_summary": summary,
        "skills": [f"Skill {index}" for index in range(18)],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
            "relevant_coursework": [f"Course {index}" for index in range(10)],
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Work bullet one.", "Work bullet two.", "Work bullet three."],
        }],
        "projects": [
            {"name": "Senior Capstone Project", "bullets": ["Capstone"] * 5},
            {"name": "Job Search Assistant", "bullets": ["JSA"] * 3},
        ],
        "certifications": ["FE Civil Exam Candidate", "Professional Engineer licensure goal"],
    }

    cleaned = generator._qa_resume_json(data, "civil engineering role")

    assert cleaned["professional_summary"].endswith(".")
    assert not cleaned["professional_summary"].endswith("Applied civil")


def test_resume_one_page_trim_preserves_high_value_ats_skills():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    generator._profile = sample_profile()
    resume = {
        "professional_summary": "Civil engineering candidate with structural project and workflow automation evidence.",
        "skills": [
            "Generic skill 1",
            "Generic skill 2",
            "Generic skill 3",
            "Generic skill 4",
            "Generic skill 5",
            "Generic skill 6",
            "Generic skill 7",
            "Generic skill 8",
            "Generic skill 9",
            "Generic skill 10",
            "Generic skill 11",
            "RAM Structural System",
            "Excel",
            "Google Sheets",
            "SQLite",
            "Python automation",
        ],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "Bachelor of Science",
            "major": "Civil Engineering Technology",
            "graduation": "May 2026",
            "relevant_coursework": ["Structural Design"] * 9,
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Supervised teams.", "Improved tracking.", "Coordinated logistics."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "bullets": ["Capstone one.", "Capstone two.", "Capstone three.", "Capstone four."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Automation Project",
                "bullets": ["Built automation.", "Maintained SQLite workflow."],
            },
        ],
        "certifications": ["FE Civil Exam Candidate"],
    }

    cleaned = generator._qa_resume_json(resume, "civil structural role using data tracking and automation")
    skills_text = " | ".join(cleaned["skills"])

    assert generator.resume_renderer_qa(cleaned)["estimated_page_count"] == 1
    assert "RAM Structural System" in skills_text
    assert "Excel" in skills_text
    assert "Google Sheets" in skills_text
    assert "SQLite" in skills_text
    assert "Python automation" in skills_text


def test_generator_falls_back_to_full_profile_when_evidence_is_sparse():
    fake = FakeLLMProvider()
    generator = DocumentGenerator(llm_provider=fake)
    generator._profile = {
        "identity": {"first_name": "James"},
        "legacy_fact": "Fallback profile fact",
    }

    result = generator.generate(make_job("General civil engineer role."))

    assert result["evidence_fallback_used"] is True
    assert "full_profile_fallback" in user_prompt(fake)
    assert "Fallback profile fact" in user_prompt(fake)


def test_resume_qa_limits_bullets_sections_and_skills():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Civil engineering candidate with project experience.",
        "skills": [f"Skill {i}" for i in range(30)],
        "education": [],
        "experience": [
            {"employer": f"Employer {i}", "title": "Lead", "bullets": [f"Work bullet {j}" for j in range(8)]}
            for i in range(4)
        ],
        "projects": [
            {"name": f"Project {i}", "role": "Student", "bullets": [f"Project bullet {j}" for j in range(8)]}
            for i in range(5)
        ],
        "activities": [],
        "certifications": [],
    }

    out = generator._qa_resume_json(data, "Structural engineering role with steel and concrete.")

    assert len(out["skills"]) == 14
    assert len(out["experience"]) == 1
    assert 1 <= len(out["experience"][0]["bullets"]) <= 3
    assert 1 <= len(out["projects"]) <= 3
    assert all(len(proj["bullets"]) <= 4 for proj in out["projects"])
    assert out["rendering_qa"]["length_risk"] in {"ok", "soft_cap_risk", "hard_cap_risk"}
    assert out["rendering_qa"]["estimated_words"] <= 950


def test_resume_qa_preserves_work_and_prefers_job_search_assistant_breadth():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Structural candidate.",
        "skills": ["RAM Structural System", "Python", "SQLite"],
        "education": [],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "dates": "September 2021 - Present",
            "location": "Lake Grove, New York",
            "bullets": ["Coordinated staffing.", "Trained employees.", "Managed operations."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "bullets": ["Capstone bullet 1", "Capstone bullet 2", "Capstone bullet 3", "Capstone bullet 4"],
            },
            {
                "name": "Steel Design Project",
                "role": "Student designer",
                "bullets": ["Duplicate steel structural design evidence."],
            },
            {
                "name": "Job Search Assistant",
                "role": "Personal software / automation project",
                "bullets": ["Built Python automation with SQLite, OpenAI, Google Sheets, and CLI workflow."],
            },
        ],
        "activities": [],
        "certifications": [],
    }

    out = generator._qa_resume_json(data, "Entry level structural engineer using RAM and steel design.")
    project_names = [project["name"] for project in out["projects"]]

    assert "Senior Capstone Project" in project_names
    assert "Job Search Assistant" in project_names
    assert "Steel Design Project" not in project_names
    assert out["experience"][0]["employer"] == "Urban Air Adventure Park"
    assert out["experience"][0]["title"] == "Event Coordination Department Head"
    assert 2 <= len(out["experience"][0]["bullets"]) <= 3
    assert len(out["projects"][0]["bullets"]) == 4
    assert out["rendering_qa"]["expansion_order"][0] == "capstone_bullets"
    if "work_experience_bullets" in out["rendering_qa"]["expansion_order"]:
        assert out["rendering_qa"]["expansion_order"].index("work_experience_bullets") > out["rendering_qa"]["expansion_order"].index("capstone_bullets")


def test_resume_qa_expands_project_evidence_before_extra_work_depth():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Civil engineering candidate.",
        "skills": ["RAM Structural System", "AutoCAD", "Steel Design"],
        "education": [],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "dates": "September 2021 - Present",
            "location": "Lake Grove, New York",
            "bullets": ["Coordinated staffing.", "Trained employees.", "Managed operations."],
        }],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "bullets": ["Capstone bullet 1", "Capstone bullet 2", "Capstone bullet 3", "Capstone bullet 4"],
            },
            {
                "name": "Stormwater Cistern Project",
                "role": "Student designer",
                "bullets": ["Calculated runoff.", "Evaluated storage and overflow options."],
            },
        ],
        "activities": [],
        "certifications": [],
    }

    out = generator._qa_resume_json(data, "Entry level civil engineer focused on structural and water resources projects.")
    order = out["rendering_qa"]["expansion_order"]

    assert out["experience"]
    assert len(out["experience"][0]["bullets"]) >= 2
    assert len(out["projects"][0]["bullets"]) == 4
    assert len(out["projects"][1]["bullets"]) == 1
    assert order.index("capstone_bullets") < order.index("academic_project_bullets")
    if "work_experience_bullets" in order:
        assert order.index("academic_project_bullets") < order.index("work_experience_bullets")


def test_resume_qa_prioritizes_job_search_assistant_before_other_academic_projects():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Civil engineering candidate.",
        "skills": ["RAM Structural System", "Python", "SQLite"],
        "education": [],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Coordinated staffing.", "Trained employees."],
        }],
        "projects": [
            {"name": "Senior Capstone Project", "role": "Structural Design Lead", "bullets": ["Capstone 1", "Capstone 2"]},
            {"name": "Bridge Replacement Project", "role": "Student designer", "bullets": ["Bridge 1"]},
            {"name": "Job Search Assistant", "role": "Software Automation Project", "bullets": ["JSA 1"]},
        ],
        "activities": [],
        "certifications": [],
    }

    out = generator._qa_resume_json(data, "Entry level structural engineer using RAM and Python automation.")
    project_names = [project["name"] for project in out["projects"]]

    assert project_names[:2] == ["Senior Capstone Project", "Job Search Assistant"]
    assert project_names.index("Job Search Assistant") < project_names.index("Bridge Replacement Project")
    assert out["experience"]


def test_coursework_trims_by_full_table_rows():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "B.S.",
            "relevant_coursework": [f"Course {idx}" for idx in range(1, 8)],
        }]
    }

    assert generator._trim_coursework_row(data) is True

    remaining = data["education"][0]["relevant_coursework"]
    assert remaining == [f"Course {idx}" for idx in range(1, 7)]


def test_length_trim_preserves_work_and_removes_third_project_before_work():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": " ".join(["summary"] * 900),
        "skills": [f"Skill {idx}" for idx in range(14)],
        "education": [{
            "institution": "Farmingdale State College",
            "degree": "B.S.",
            "relevant_coursework": [f"Course {idx}" for idx in range(1, 10)],
        }],
        "experience": [{
            "employer": "Urban Air Adventure Park",
            "title": "Event Coordination Department Head",
            "bullets": ["Coordinated staffing.", "Trained employees."],
        }],
        "projects": [
            {"name": "Senior Capstone Project", "role": "Structural Design Lead", "bullets": ["Capstone"] * 5},
            {"name": "Job Search Assistant", "role": "Software Automation Project", "bullets": ["JSA"] * 3},
            {"name": "Bridge Replacement Project", "role": "Student designer", "bullets": ["Bridge"] * 4},
        ],
        "activities": [],
        "certifications": [],
    }
    warnings = []

    generator._trim_resume_to_length(data, warnings)

    assert data["experience"]
    assert all(project["name"] != "Bridge Replacement Project" for project in data["projects"])
    assert [project["name"] for project in data["projects"]] == ["Senior Capstone Project", "Job Search Assistant"]


def test_resume_qa_removes_question_mark_replacement_characters():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Candidate ? structural.",
        "skills": ["RAM Structural System ? Project coordination"],
        "education": [],
        "experience": [],
        "projects": [],
        "activities": [],
        "certifications": [],
    }

    out = generator._qa_resume_json(data, "Structural engineer.")

    assert "?" not in out["professional_summary"]
    assert "?" not in " ".join(out["skills"])


def test_relevant_coursework_is_conditional():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "Candidate.",
        "skills": [],
        "education": [{
            "degree": "B.S.",
            "major": "Civil Engineering Technology",
            "institution": "Farmingdale State College",
            "graduation": "May 2026",
            "relevant_coursework": ["Structural Design", "Reinforced Concrete Design"],
        }],
        "experience": [],
        "projects": [],
        "activities": [],
        "certifications": [],
    }

    relevant = generator._qa_resume_json(data, "Entry level structural engineer working with concrete.")
    irrelevant = generator._qa_resume_json(data, "Customer success role focused on sales outreach.")

    assert relevant["education"][0]["relevant_coursework"]
    assert "relevant_coursework" not in irrelevant["education"][0]
    assert "coursework omitted" in " ".join(irrelevant["rendering_qa"]["warnings"])


def test_resume_qa_removes_blank_sections_and_items():
    generator = DocumentGenerator(llm_provider=FakeLLMProvider())
    data = {
        "professional_summary": "  ",
        "skills": ["", "RAM"],
        "education": [{"degree": "", "institution": ""}],
        "experience": [{"employer": "", "title": "", "bullets": [""]}],
        "projects": [{"name": "", "role": "", "bullets": [""]}],
        "activities": [{"organization": "", "role": ""}],
        "certifications": ["", "FE Civil Exam Candidate"],
    }

    out = generator._qa_resume_json(data, "Structural engineer.")

    assert out["professional_summary"] == ""
    assert out["skills"] == ["RAM"]
    assert out["education"] == []
    assert out["experience"] == []
    assert out["projects"] == []
    assert out["activities"] == []
    assert out["certifications"] == ["FE Civil Exam Candidate"]


def test_generation_and_grading_do_not_import_openai_directly():
    generator_source = Path("job_search/generation/generator.py").read_text(encoding="utf-8")
    grader_source = Path("job_search/grading/grader.py").read_text(encoding="utf-8")

    assert "from openai import OpenAI" not in generator_source
    assert "from openai import OpenAI" not in grader_source

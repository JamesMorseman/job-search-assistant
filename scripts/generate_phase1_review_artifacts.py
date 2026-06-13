"""Generate local Phase 1 review artifacts for final visual inspection.

This script is intentionally standalone and does not participate in runtime
generation workflows.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

import yaml
from docx import Document

from job_search.db import get_db
from job_search.generation.audit import audit_generated_documents, cover_docx_formatting_metadata
from job_search.generation.generator import DocumentGenerator
from job_search.models import ATSType, CanonicalJob


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts" / "phase1_review"
REAL_JOB_OUT_DIR = ROOT / "artifacts" / "phase1_real_job"


def sample_resume_json() -> dict:
    return {
        "professional_summary": (
            "Entry-level civil engineering candidate focused on structural "
            "analysis, project documentation, and practical engineering tools."
        ),
        "skills": [
            "RAM Structural System",
            "AutoCAD",
            "Structural Analysis",
            "Steel Design",
            "Python",
            "SQLite",
        ],
        "education": [
            {
                "degree": "B.S.",
                "major": "Civil Engineering Technology",
                "institution": "Farmingdale State College",
                "location": "Farmingdale, New York",
                "graduation": "Graduated May 2026",
                "honors": ["Dean's List - Spring 2026"],
                "relevant_coursework": [
                    "Structural Design",
                    "Reinforced Concrete Design",
                    "Steel Design",
                    "Soils and Foundations",
                    "Civil Engineering Materials",
                    "Hydraulics",
                ],
            }
        ],
        "projects": [
            {
                "name": "Senior Capstone Project",
                "role": "Structural Design Lead",
                "organization": "Farmingdale State College",
                "date": "Spring 2026",
                "bullets": [
                    "Built and reviewed structural models, load paths, and member sizing alternatives."
                ],
            },
            {
                "name": "Job Search Assistant",
                "role": "Software Automation Project",
                "date": "2026",
                "bullets": [
                    "Built Python automation with SQLite-backed tracking, document workflows, and tests."
                ],
            },
        ],
        "experience": [
            {
                "employer": "Urban Air Adventure Park",
                "title": "Event Coordination Department Head",
                "dates": "September 2021 - Present",
                "location": "Lake Grove, New York",
                "bullets": [
                    "Coordinated staffing, logistics, training, and event execution for high-volume events."
                ],
            }
        ],
        "activities": [],
        "certifications": ["FE Civil Exam Candidate (August 2026)"],
    }


def sample_cover_json(generator: DocumentGenerator) -> dict:
    return generator._qa_cover_letter_json({
        "salutation": "Dear Hiring Manager,",
        "body_paragraphs": [
            (
                "I am interested in the Civil Engineer role at Acme "
                "[Infrastructure Group] because it aligns with my structural "
                "coursework, project documentation experience, and interest in "
                "practical engineering delivery."
            ),
            (
                "My senior capstone work involved structural modeling, load-path "
                "review, and engineering documentation, while related coursework "
                "strengthened my foundation in steel, reinforced concrete, and "
                "soils concepts."
            ),
            (
                "I also bring operations leadership and automation initiative "
                "through coordinating teams and building the Job Search Assistant "
                "to organize job data, document generation, and follow-up workflows."
            ),
        ],
        "closing": "Sincerely, James Morseman",
    })


def _load_profile() -> dict:
    return yaml.safe_load((ROOT / "profile" / "james_profile.yaml").read_text(encoding="utf-8"))


def _job_from_row(row) -> CanonicalJob:
    return CanonicalJob(
        source=row["source"],
        source_job_id=row["source_job_id"],
        firm_id=row["firm_id"],
        company=row["company"],
        title=row["title"],
        location_city=row["location_city"],
        location_state=row["location_state"],
        description_normalized=row["description_normalized"],
        description_raw=row["description_raw"],
        apply_url=row["apply_url"],
        ats_type=ATSType(row["ats_type"] or "unknown"),
    )


def _docx_text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs if paragraph.text)


def generate_real_job_artifacts(job_id: str) -> None:
    REAL_JOB_OUT_DIR.mkdir(parents=True, exist_ok=True)

    with get_db() as db:
        row = db.execute("SELECT * FROM jobs WHERE canonical_job_id = ?", (job_id,)).fetchone()
    if row is None:
        raise ValueError(f"Job not found: {job_id}")

    profile = _load_profile()
    job = _job_from_row(row)
    generator = DocumentGenerator()
    generator._profile = profile
    result = generator.generate(job)

    safe_company = "".join(c for c in job.company if c.isalnum() or c in " -_")[:30]
    safe_title = "".join(c for c in job.title if c.isalnum() or c in " -_")[:30]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = f"{date.today().isoformat()}_{timestamp}_{safe_company}_{safe_title}".replace(" ", "_")
    resume_path = REAL_JOB_OUT_DIR / f"{prefix}_resume.docx"
    cover_path = REAL_JOB_OUT_DIR / f"{prefix}_cover_letter.docx"
    audit_path = REAL_JOB_OUT_DIR / f"{prefix}_audit.json"
    summary_path = REAL_JOB_OUT_DIR / f"{prefix}_render_summary.txt"

    rendered_resume_json, _ = generator.prepare_resume_for_rendering(result["resume_json"])
    generator.save_docx(rendered_resume_json, str(resume_path))
    generator.save_cover_docx(result["cover_letter_json"], str(cover_path), job=job, today=date.today().isoformat())
    audit = audit_generated_documents(
        resume_json=rendered_resume_json,
        cover_letter_json=result["cover_letter_json"],
        profile=profile,
        job=job,
        resume_text=_docx_text(resume_path),
        cover_letter_text=_docx_text(cover_path),
        cover_formatting_metadata=cover_docx_formatting_metadata(str(cover_path)),
    )
    resume_qa = generator.resume_renderer_qa(rendered_resume_json)
    audit_payload = {
        "job_id": job_id,
        "company": job.company,
        "title": job.title,
        "resume_renderer_qa": resume_qa,
        "document_audit": audit,
    }
    audit_path.write_text(json.dumps(audit_payload, indent=2), encoding="utf-8")

    resume_doc = Document(resume_path)
    cover_doc = Document(cover_path)
    skill_groups = generator._group_skills(  # noqa: SLF001 - review artifact diagnostic.
        rendered_resume_json.get("skills", []),
        separate_programming_data=generator._is_software_data_heavy_resume(rendered_resume_json),  # noqa: SLF001
        resume_json=rendered_resume_json,
    )
    skill_lines = [f"{label}: {' | '.join(values)}" for label, values in skill_groups.items()]
    summary = [
        "Phase 1 Real-Job Artifact Summary",
        "",
        f"Job ID: {job_id}",
        f"Company: {job.company}",
        f"Title: {job.title}",
        f"Resume: {resume_path}",
        f"Cover letter: {cover_path}",
        f"Audit: {audit_path}",
        "",
        "Resume header:",
        *[p.text for p in resume_doc.paragraphs[:3]],
        "",
        "Resume technical skills:",
        *skill_lines,
        "",
        f"Resume QA warnings: {resume_qa['warnings']}",
        f"Document audit status: {audit['status']}",
        "",
        "Cover-letter leading block:",
        *[p.text for p in cover_doc.paragraphs[:8] if p.text],
        "",
        "Cover-letter closing block:",
        *[p.text for p in cover_doc.paragraphs[-3:] if p.text],
        "",
        "Cover-letter audit warnings:",
        *[warning["code"] for warning in audit["cover_letter"]["warnings"]],
    ]
    summary_path.write_text("\n".join(summary), encoding="utf-8")

    print(resume_path)
    print(cover_path)
    print(audit_path)
    print(summary_path)


def generate_sample_artifacts() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    profile = _load_profile()
    generator = DocumentGenerator(llm_provider=object())
    generator._profile = profile

    resume_path = OUT_DIR / "phase1_review_resume.docx"
    cover_path = OUT_DIR / "phase1_review_cover_letter.docx"
    summary_path = OUT_DIR / "phase1_review_render_summary.txt"

    generator.save_docx(sample_resume_json(), str(resume_path))
    generator.save_cover_docx(
        sample_cover_json(generator),
        str(cover_path),
        job=CanonicalJob(
            source="review",
            source_job_id="phase1-review",
            company="Acme [Infrastructure Group]",
            title="Civil Engineer",
            location_city="New York",
            location_state="NY",
            description_normalized="Civil engineering role with structural and documentation work.",
        ),
        today="2026-06-12",
        include_location=True,
        include_linkedin=False,
    )

    resume_doc = Document(resume_path)
    cover_doc = Document(cover_path)
    summary = [
        "Phase 1 Review Artifact Summary",
        "",
        f"Resume: {resume_path}",
        f"Cover letter: {cover_path}",
        "",
        "Resume header:",
        *[p.text for p in resume_doc.paragraphs[:3]],
        "",
        "Cover-letter leading block:",
        *[p.text for p in cover_doc.paragraphs[:8] if p.text],
        "",
        "Cover-letter closing block:",
        *[p.text for p in cover_doc.paragraphs[-3:] if p.text],
    ]
    summary_path.write_text("\n".join(summary), encoding="utf-8")

    print(resume_path)
    print(cover_path)
    print(summary_path)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Generate local Phase 1 review artifacts.")
    parser.add_argument("--job-id", help="Generate review artifacts from a real SQLite job ID.")
    args = parser.parse_args([] if argv is None else argv)
    if args.job_id:
        generate_real_job_artifacts(args.job_id)
    else:
        generate_sample_artifacts()


if __name__ == "__main__":
    main(sys.argv[1:])

"""Deterministic Phase 1 document-generation audit checks."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from enum import Enum
from typing import Any, Literal

from docx import Document as DocxDocument

from job_search.models import CanonicalJob

DocumentType = Literal["resume", "cover_letter", "document_set"]


class AuditStatus(str, Enum):
    PASS = "PASS"
    FAIL = "FAIL"


class AuditSeverity(str, Enum):
    CRITICAL = "critical"
    MAJOR = "major"
    MINOR = "minor"


TEMPLATE_PLACEHOLDER_RE = re.compile(r"\b[\w-]*placeholder[\w-]*\b", re.IGNORECASE)
LEFTOVER_BOILERPLATE_RE = re.compile(
    r"\b(?:lorem\s+ipsum|TODO|FIXME|TBD|XXX)\b", re.IGNORECASE
)
TEMPLATE_MARKER_LABELS = (
    "company",
    "role",
    "title",
    "closing",
    "signature",
    "salutation",
    "body",
    "body_paragraph",
    "body_paragraphs",
    "hiring_manager",
    "date",
    "location",
    "candidate",
    "name",
    "email",
    "phone",
    "linkedin",
    "github",
)
TEMPLATE_MARKER_LABEL_RE = "|".join(TEMPLATE_MARKER_LABELS)
UNRESOLVED_TEMPLATE_MARKER_RE = re.compile(
    rf"""
    \{{\{{\s*[\w.-]+\s*\}}\}}
    |\$\{{\s*[\w.-]+\s*\}}
    |\{{\s*(?:{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*\}}
    |\[\s*(?:{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*\]
    |<\s*(?:placeholder|{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*>
    |(?:^|\n)\s*(?:{TEMPLATE_MARKER_LABEL_RE})\s*(?:\n|$)
    |(?:^|[\s.])[\[\{{]\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)


class DocumentAuditFailure(ValueError):
    """Raised when a generated document fails the Phase 1 audit gate."""

    def __init__(self, result: "DocumentAuditResult"):
        self.result = result
        super().__init__(result.failure_summary())


@dataclass(frozen=True)
class DocumentAuditCheck:
    code: str
    severity: AuditSeverity
    requirement: str
    explanation: str
    recommended_fix: str
    passed: bool = True

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["severity"] = self.severity.value
        return data


@dataclass(frozen=True)
class DocumentAuditResult:
    document_type: DocumentType
    checks: tuple[DocumentAuditCheck, ...] = field(default_factory=tuple)

    @property
    def failures(self) -> tuple[DocumentAuditCheck, ...]:
        return tuple(
            check
            for check in self.checks
            if not check.passed and check.severity != AuditSeverity.MINOR
        )

    @property
    def warnings(self) -> tuple[DocumentAuditCheck, ...]:
        return tuple(
            check
            for check in self.checks
            if not check.passed and check.severity == AuditSeverity.MINOR
        )

    @property
    def status(self) -> AuditStatus:
        return AuditStatus.FAIL if self.failures else AuditStatus.PASS

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_type": self.document_type,
            "status": self.status.value,
            "checks": [check.to_dict() for check in self.checks],
            "failures": [check.to_dict() for check in self.failures],
            "warnings": [check.to_dict() for check in self.warnings],
        }

    def failure_summary(self) -> str:
        if not self.failures:
            return f"{self.document_type} audit passed."
        codes = ", ".join(check.code for check in self.failures)
        return f"{self.document_type} audit failed: {codes}"

    def raise_for_failure(self) -> None:
        if self.failures:
            raise DocumentAuditFailure(self)


def audit_generated_documents(
    *,
    resume_json: dict,
    cover_letter_json: dict,
    profile: dict,
    job: CanonicalJob | None = None,
    resume_text: str = "",
    cover_letter_text: str = "",
    cover_formatting_metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Audit a generated resume and cover letter as a document set."""
    resume = audit_resume(resume_json=resume_json, profile=profile, resume_text=resume_text)
    cover = audit_cover_letter(
        cover_json=cover_letter_json,
        profile=profile,
        job=job,
        cover_text=cover_letter_text,
        formatting_metadata=cover_formatting_metadata,
    )
    set_result = DocumentAuditResult(
        document_type="document_set",
        checks=resume.checks + cover.checks,
    )
    return {
        "status": set_result.status.value,
        "resume": resume.to_dict(),
        "cover_letter": cover.to_dict(),
        "failures": [check.to_dict() for check in set_result.failures],
        "warnings": [check.to_dict() for check in set_result.warnings],
    }


def audit_resume(*, resume_json: dict, profile: dict, resume_text: str = "") -> DocumentAuditResult:
    identity = profile.get("identity", {}) if isinstance(profile, dict) else {}
    header_lines = _resume_header_lines(identity)
    text_parts = [resume_text, _json_text(resume_json), "\n".join(header_lines)]
    full_text = "\n".join(part for part in text_parts if part)
    projects = resume_json.get("projects", []) if isinstance(resume_json, dict) else []
    skills = resume_json.get("skills", []) if isinstance(resume_json, dict) else []
    estimated_page_count = _resume_estimated_page_count(resume_json, full_text)
    checks = [
        _check(bool(header_lines and header_lines[0]), "MISSING_NAME", AuditSeverity.CRITICAL,
               "Resume header must include candidate name.", "Generated resume header has no name.",
               "Verify profile identity first_name/last_name and resume header renderer."),
        _check(bool(identity.get("phone")), "MISSING_PHONE", AuditSeverity.CRITICAL,
               "Approved resume header must include phone.", "Profile or rendered header is missing phone.",
               "Verify profile.identity.phone and resume header renderer."),
        _check(bool(identity.get("email")), "MISSING_EMAIL", AuditSeverity.CRITICAL,
               "Approved resume header must include email.", "Profile or rendered header is missing email.",
               "Verify profile.identity.email and resume header renderer."),
        _check(bool(identity.get("linkedin_url")), "MISSING_LINKEDIN", AuditSeverity.CRITICAL,
               "Approved resume header must include LinkedIn.", "Profile or rendered header is missing LinkedIn URL.",
               "Verify profile.identity.linkedin_url and resume header renderer."),
        _check(bool(identity.get("github_url")), "MISSING_GITHUB", AuditSeverity.CRITICAL,
               "Approved resume header must include LinkedIn and GitHub.", "Profile or rendered header is missing GitHub URL.",
               "Verify profile.identity.github_url and resume header renderer."),
        _check(_resume_header_is_approved(header_lines, identity), "MALFORMED_HEADER", AuditSeverity.CRITICAL,
               "Resume header must follow Name / Phone | Email / LinkedIn | GitHub.",
               "Generated resume header does not match the approved three-line layout.",
               "Render resume header with phone/email first and LinkedIn/GitHub second."),
        _check(bool(_clean(resume_json.get("professional_summary", ""))), "MISSING_PROFESSIONAL_SUMMARY", AuditSeverity.CRITICAL,
               "Professional Summary must be present.", "Resume JSON has no Professional Summary.",
               "Regenerate or repair resume_json.professional_summary."),
        _check(bool(resume_json.get("education")), "MISSING_EDUCATION", AuditSeverity.CRITICAL,
               "Education section must be present.", "Resume JSON has no Education entries.",
               "Regenerate or restore profile-grounded education."),
        _check(bool(projects), "MISSING_ENGINEERING_EXPERIENCE", AuditSeverity.CRITICAL,
               "Engineering Experience must be present.", "Resume JSON has no project/engineering experience entries.",
               "Regenerate with capstone and engineering project evidence."),
        _check(bool(skills), "MISSING_TECHNICAL_SKILLS", AuditSeverity.CRITICAL,
               "Technical Skills must be present.", "Resume JSON has no technical skills.",
               "Regenerate or restore profile-grounded technical skills."),
        _check(_has_capstone_or_baldwin(full_text), "MISSING_CAPSTONE_EVIDENCE", AuditSeverity.CRITICAL,
               "Resume must include capstone/Baldwin High School evidence.", "No capstone or Baldwin High School evidence was detected.",
               "Include the flagship capstone project or verified Baldwin High School project evidence."),
        _check(not contains_unresolved_template_marker(full_text), "PLACEHOLDER_LEAKAGE", AuditSeverity.CRITICAL,
               "Generated resume must not contain unresolved placeholders or template markers.",
               "Resume contains unresolved placeholder/schema/template text.",
               "Regenerate or strip unresolved template markers before rendering."),
        _check(estimated_page_count <= 1, "PAGE_COUNT_EXCEEDED", AuditSeverity.CRITICAL,
               "Generated resume must not exceed the approved one-page maximum.",
               "Generated resume exceeds approved one-page limit.",
               "Trim excess coursework, lower-priority bullets, skills, or professional-development detail before rendering."),
        _check(not _has_duplicate_headings(full_text), "DUPLICATE_SECTION_HEADING", AuditSeverity.MAJOR,
               "Resume section headings should not duplicate.", "Duplicate section headings were detected.",
               "Deduplicate generated resume sections before rendering."),
        _check(_skills_have_engineering_keywords(skills), "WEAK_TECHNICAL_SKILLS", AuditSeverity.MAJOR,
               "Technical Skills should include engineering keywords.", "Technical skills appear sparse or not engineering-oriented.",
               "Include role-relevant engineering tools, methods, and standards."),
        _check(_has_job_search_assistant_if_profile_has_it(profile, full_text), "MISSING_JSA_EVIDENCE", AuditSeverity.MAJOR,
               "Job Search Assistant evidence should appear when profile contains it and generation uses it.",
               "Profile contains Job Search Assistant evidence but resume text does not.",
               "Include Job Search Assistant when selected/relevant, framed as automation/data/workflow evidence."),
        _check(len(projects) == 0 or len(projects[0].get("bullets", [])) >= 1, "SPARSE_CAPSTONE_BULLETS", AuditSeverity.MAJOR,
               "Capstone project should have supporting bullets.", "First engineering project has no supporting bullet evidence.",
               "Add concise capstone/project bullets grounded in profile evidence."),
        _check(not _has_repeated_bullet_text(projects), "DUPLICATE_BULLET_TEXT", AuditSeverity.MAJOR,
               "Project/engineering experience bullets should not repeat identical text.",
               "Two or more project bullets contain identical text.",
               "Regenerate or rewrite duplicated bullets so each conveys distinct evidence."),
        _check(_resume_utilization_ok(full_text), "RESUME_UNDERUTILIZED", AuditSeverity.MINOR,
               "Resume should use available page capacity reasonably.", "Resume appears short for a one-page target.",
               "Consider adding relevant coursework, skills, or project detail if space allows."),
    ]
    return DocumentAuditResult("resume", tuple(checks))


def audit_cover_letter(
    *,
    cover_json: dict,
    profile: dict,
    job: CanonicalJob | None = None,
    cover_text: str = "",
    formatting_metadata: dict[str, Any] | None = None,
) -> DocumentAuditResult:
    identity = profile.get("identity", {}) if isinstance(profile, dict) else {}
    body = cover_json.get("body_paragraphs", []) if isinstance(cover_json, dict) else []
    paragraphs = [str(p).strip() for p in body if str(p).strip()] if isinstance(body, list) else []
    closing = str(cover_json.get("closing", "") or "") if isinstance(cover_json, dict) else ""
    salutation = _clean(cover_json.get("salutation", "")) if isinstance(cover_json, dict) else ""
    full_text = "\n".join([cover_text, salutation, *paragraphs, closing])
    checks = [
        _check(_has_date(full_text), "MISSING_DATE", AuditSeverity.CRITICAL,
               "Cover letter must include a date.", "No rendered date was detected in the cover letter.",
               "Render the cover letter with a date before upload/output."),
        _check(_company_present(full_text, job), "MISSING_COMPANY", AuditSeverity.CRITICAL,
               "Cover letter must include employer name when available.", "Target employer name was not detected in the cover letter.",
               "Render the company/employer block and include role/company specificity."),
        _check(bool(salutation), "MISSING_GREETING", AuditSeverity.CRITICAL,
               "Cover letter greeting must be present.", "Cover letter salutation is missing.",
               "Use a verified named greeting or Dear Hiring Manager."),
        _check(len(paragraphs) > 1, "SINGLE_PARAGRAPH_COLLAPSE", AuditSeverity.CRITICAL,
               "Cover letter must use multiple body paragraphs.", "Body appears collapsed into one paragraph.",
               "Regenerate with separated body_paragraphs."),
        _check(len(paragraphs) >= 3, "TOO_FEW_BODY_PARAGRAPHS", AuditSeverity.MAJOR,
               "Cover letter architecture requires 3 body paragraphs by default.",
               "Cover letter has fewer than 3 body paragraphs.",
               "Regenerate with role/company, engineering evidence, and leadership/automation paragraphs."),
        _check(_has_engineering_evidence(full_text), "MISSING_ENGINEERING_EVIDENCE", AuditSeverity.CRITICAL,
               "Cover letter must include engineering evidence.", "No engineering/project evidence was detected.",
               "Include capstone, coursework, technical tools, or engineering project evidence."),
        _check(_has_capstone_or_baldwin(full_text), "MISSING_FLAGSHIP_EVIDENCE", AuditSeverity.CRITICAL,
               "Cover letter must include capstone or equivalent flagship evidence.",
               "No capstone/equivalent flagship evidence was detected.",
               "Mention capstone or another verified flagship engineering project."),
        _check(_closing_present(closing), "MISSING_CLOSING", AuditSeverity.CRITICAL,
               "Cover letter closing must be present.", "Professional closing is missing.",
               "Use a distinct professional closing such as Sincerely,."),
        _check(_signature_present(closing, identity), "MISSING_SIGNATURE", AuditSeverity.CRITICAL,
               "Cover letter signature must be present.", "Candidate signature is missing.",
               "Include James Morseman on a separate signature line."),
        _check(_closing_signature_separate(closing), "MERGED_CLOSING_SIGNATURE", AuditSeverity.CRITICAL,
               "Closing and signature must be separated from body text.", "Closing/signature are missing a line break or appear merged.",
               "Render closing and signature as separate paragraphs/lines."),
        _check(not contains_unresolved_template_marker(full_text), "PLACEHOLDER_LEAKAGE", AuditSeverity.CRITICAL,
               "Generated cover letter must not contain unresolved placeholders or template markers.",
               "Cover letter contains unresolved placeholder/schema/template text.",
               "Regenerate or strip unresolved template markers before rendering."),
        _check(not _has_malformed_closing_fragment(full_text), "MALFORMED_CLOSING_FRAGMENT", AuditSeverity.CRITICAL,
               "Cover letter must not contain malformed leaked fragments like closing. [.",
               "Malformed closing/template fragment was detected.",
               "Regenerate and validate raw cover-letter body before rendering."),
        _check(_contact_block_possible(identity), "MISSING_CONTACT_BLOCK", AuditSeverity.MAJOR,
               "Cover-letter header should include stable contact information.",
               "Profile lacks phone/email needed for cover-letter contact block.",
               "Verify profile identity phone/email before generation."),
        _check(_company_location_block_present(full_text, job), "MISSING_COMPANY_LOCATION_BLOCK", AuditSeverity.MAJOR,
               "Cover letter should include company/location block when available.",
               "Company/location block was not detected in rendered cover letter.",
               "Render company and location lines when the job has company/location data."),
        _check(_job_specificity_ok(full_text, job), "GENERIC_COVER_LETTER", AuditSeverity.MAJOR,
               "Cover letter should be specific to role/company.", "Role/company specificity is weak.",
               "Mention the target company and role naturally."),
        _check(_has_leadership_or_jsa_if_expected(profile, full_text), "MISSING_LEADERSHIP_OR_JSA", AuditSeverity.MAJOR,
               "Leadership or Job Search Assistant evidence should appear when expected.",
               "Profile contains leadership/JSA evidence but cover letter does not mention either.",
               "Include relevant leadership or automation/workflow evidence."),
        _check(False, "COVER_LOCATION_REVIEW", AuditSeverity.MINOR,
               "Cover-letter location may be configurable.", "Location inclusion may need human review.",
               "Keep include_location configurable for final cover-letter rendering."),
    ]
    if formatting_metadata is not None:
        checks.extend(_cover_formatting_checks(formatting_metadata))
    return DocumentAuditResult("cover_letter", tuple(checks))


def contains_unresolved_template_marker(value: str) -> bool:
    text = value or ""
    return bool(
        TEMPLATE_PLACEHOLDER_RE.search(text)
        or UNRESOLVED_TEMPLATE_MARKER_RE.search(text)
        or _has_malformed_closing_fragment(text)
        or _has_leftover_boilerplate(text)
    )


def cover_docx_formatting_metadata(path: str) -> dict[str, Any]:
    """Inspect rendered DOCX paragraph formatting used by the cover-letter audit gate."""
    doc = DocxDocument(path)
    paragraphs = [p for p in doc.paragraphs if _clean(p.text)]
    texts = [_clean(p.text) for p in paragraphs]
    greeting_index = _first_index(texts, lambda text: text.lower().startswith("dear "))
    closing_index = _first_index(texts, lambda text: _closing_present(text))
    signature_index = _first_index_after(
        texts,
        closing_index,
        lambda text: "james morseman" in text.lower(),
    )
    date_index = _first_index(texts, _has_date)
    body_paragraphs = (
        paragraphs[greeting_index + 1:closing_index]
        if greeting_index >= 0 and closing_index > greeting_index
        else []
    )
    return {
        "paragraph_count": len(paragraphs),
        "line_spacing_configured": bool(paragraphs)
        and all(p.paragraph_format.line_spacing is not None for p in paragraphs),
        "body_paragraph_spacing_configured": bool(body_paragraphs)
        and all(_points(p.paragraph_format.space_after) > 0 for p in body_paragraphs),
        "closing_signature_separation_configured": (
            closing_index >= 0
            and signature_index > closing_index
            and _points(paragraphs[closing_index].paragraph_format.space_before) > 0
            and paragraphs[closing_index].paragraph_format.line_spacing is not None
            and paragraphs[signature_index].paragraph_format.line_spacing is not None
        ),
        "distinct_contact_date_body_blocks": (
            len(paragraphs) >= 8
            and date_index >= 2
            and greeting_index > date_index
            and closing_index > greeting_index
            and signature_index > closing_index
        ),
    }


def audit_report_markdown(report: dict[str, Any]) -> str:
    lines = [f"# Document Audit: {report['status']}", ""]
    for key, title in (("resume", "Resume"), ("cover_letter", "Cover Letter")):
        section = report[key]
        lines += [f"## {title}: {section['status']}", ""]
        failures = section.get("failures", [])
        if failures:
            for failure in failures:
                lines.append(
                    f"- {failure['code']} ({failure['severity']}): "
                    f"{failure['explanation']} Fix: {failure['recommended_fix']}"
                )
        else:
            lines.append("- No failures.")
        lines.append("")
    return "\n".join(lines)


def _check(
    passed: bool,
    code: str,
    severity: AuditSeverity,
    requirement: str,
    explanation: str,
    recommended_fix: str,
) -> DocumentAuditCheck:
    return DocumentAuditCheck(
        code=code,
        severity=severity,
        requirement=requirement,
        explanation=explanation,
        recommended_fix=recommended_fix,
        passed=passed,
    )


def _cover_formatting_checks(metadata: dict[str, Any]) -> list[DocumentAuditCheck]:
    return [
        _check(bool(metadata.get("body_paragraph_spacing_configured")), "COVER_BODY_SPACING_NOT_CONFIGURED", AuditSeverity.MINOR,
               "Cover-letter body paragraphs must have explicit paragraph spacing.",
               "Rendered DOCX body paragraphs do not show explicit spacing configuration.",
               "Apply explicit paragraph_format.space_after to cover-letter body paragraphs."),
        _check(bool(metadata.get("line_spacing_configured")), "COVER_LINE_SPACING_NOT_CONFIGURED", AuditSeverity.MINOR,
               "Cover-letter paragraphs must have explicit readable line spacing.",
               "Rendered DOCX paragraphs do not all show explicit line-spacing configuration.",
               "Apply explicit paragraph_format.line_spacing to rendered cover-letter paragraphs."),
        _check(bool(metadata.get("closing_signature_separation_configured")), "COVER_CLOSING_SIGNATURE_SPACING_NOT_CONFIGURED", AuditSeverity.MINOR,
               "Closing and signature must be visually separated from the body and each other.",
               "Rendered DOCX does not show configured closing/signature paragraph separation.",
               "Render closing and signature as distinct paragraphs with explicit before/after spacing."),
        _check(bool(metadata.get("distinct_contact_date_body_blocks")), "COVER_BLOCK_STRUCTURE_NOT_DISTINCT", AuditSeverity.MINOR,
               "Contact, date, employer, greeting, body, closing, and signature should use distinct paragraph blocks.",
               "Rendered DOCX paragraph structure does not clearly separate cover-letter blocks.",
               "Render each business-letter block as its own paragraph with explicit spacing."),
    ]


def _resume_header_lines(identity: dict) -> tuple[str, str, str]:
    first = _clean(identity.get("first_name", ""))
    last = _clean(identity.get("last_name", ""))
    name = " ".join(part for part in (first, last) if part)
    primary = " | ".join(part for part in (_clean(identity.get("phone", "")), _clean(identity.get("email", ""))) if part)
    links = " | ".join(part for part in (_clean(identity.get("linkedin_url", "")), _clean(identity.get("github_url", ""))) if part)
    return name, primary, links


def _resume_header_is_approved(header_lines: tuple[str, str, str], identity: dict) -> bool:
    name, primary, links = header_lines
    return bool(
        name
        and primary == f"{_clean(identity.get('phone', ''))} | {_clean(identity.get('email', ''))}"
        and links == f"{_clean(identity.get('linkedin_url', ''))} | {_clean(identity.get('github_url', ''))}"
    )


def _json_text(value: Any) -> str:
    if isinstance(value, dict):
        return " ".join(_json_text(v) for v in value.values())
    if isinstance(value, list):
        return " ".join(_json_text(v) for v in value)
    return str(value or "")


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _first_index(values: list[str], predicate) -> int:
    for index, value in enumerate(values):
        if predicate(value):
            return index
    return -1


def _first_index_after(values: list[str], start_index: int, predicate) -> int:
    if start_index < 0:
        return -1
    for index, value in enumerate(values[start_index + 1:], start=start_index + 1):
        if predicate(value):
            return index
    return -1


def _points(value: Any) -> float:
    if value is None:
        return 0.0
    if hasattr(value, "pt"):
        return float(value.pt)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _has_leftover_boilerplate(text: str) -> bool:
    return bool(LEFTOVER_BOILERPLATE_RE.search(text or ""))


def _has_repeated_bullet_text(projects: list[Any]) -> bool:
    bullets = []
    for project in projects or []:
        if not isinstance(project, dict):
            continue
        for bullet in project.get("bullets", []) or []:
            cleaned = _clean(bullet).lower()
            if cleaned:
                bullets.append(cleaned)
    if len(bullets) < 2:
        return False
    return len(set(bullets)) < len(bullets)


def _has_capstone_or_baldwin(text: str) -> bool:
    lowered = text.lower()
    return "capstone" in lowered or "baldwin high school" in lowered


def _has_engineering_evidence(text: str) -> bool:
    lowered = text.lower()
    terms = ("engineering", "structural", "civil", "steel", "concrete", "ram", "capstone", "coursework")
    return any(term in lowered for term in terms)


def _skills_have_engineering_keywords(skills: list[Any]) -> bool:
    text = _json_text(skills).lower()
    terms = ("structural", "civil", "steel", "concrete", "ram", "autocad", "revit", "analysis", "design")
    return any(term in text for term in terms)


def _has_duplicate_headings(text: str) -> bool:
    headings = [
        "professional summary",
        "education",
        "engineering experience",
        "work experience",
        "technical skills",
        "professional development",
    ]
    lines = [_clean(line).lower() for line in (text or "").splitlines() if _clean(line)]
    return any(lines.count(heading) > 1 for heading in headings)


def _resume_utilization_ok(text: str) -> bool:
    return len(text.split()) >= 250


def _resume_estimated_page_count(resume_json: dict, text: str) -> int:
    qa = resume_json.get("rendering_qa", {}) if isinstance(resume_json, dict) else {}
    try:
        return int(qa.get("estimated_page_count") or 1)
    except (TypeError, ValueError):
        return 2 if len((text or "").split()) > 700 else 1


def _has_job_search_assistant_if_profile_has_it(profile: dict, text: str) -> bool:
    profile_text = _json_text(profile).lower()
    if "job search assistant" not in profile_text:
        return True
    return "job search assistant" in text.lower()


def _closing_present(closing: str) -> bool:
    return any(term in closing.lower() for term in ("sincerely", "regards", "respectfully"))


def _signature_present(closing: str, identity: dict) -> bool:
    first = _clean(identity.get("first_name", "")) or "James"
    last = _clean(identity.get("last_name", "")) or "Morseman"
    return f"{first} {last}".lower() in closing.lower()


def _closing_signature_separate(closing: str) -> bool:
    return "\n" in str(closing or "")


def _has_malformed_closing_fragment(text: str) -> bool:
    return bool(
        re.search(r"\bclosing\.\s*[\[\{]\s*$", text or "", flags=re.IGNORECASE | re.MULTILINE)
        or re.search(r"(?m)^\s*closing\s*$", text or "", flags=re.IGNORECASE)
    )


def _contact_block_possible(identity: dict) -> bool:
    return bool(identity.get("phone") and identity.get("email"))


def _has_date(text: str) -> bool:
    return bool(
        re.search(r"\b20\d{2}-\d{2}-\d{2}\b", text or "")
        or re.search(
            r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+20\d{2}\b",
            text or "",
        )
    )


def _company_present(text: str, job: CanonicalJob | None) -> bool:
    if job is None or not _clean(getattr(job, "company", "")):
        return True
    return _clean(getattr(job, "company", "")).lower() in (text or "").lower()


def _company_location_block_present(text: str, job: CanonicalJob | None) -> bool:
    if job is None:
        return True
    if not _company_present(text, job):
        return False
    location = ", ".join(part for part in [getattr(job, "location_city", None), getattr(job, "location_state", None)] if part)
    if not location:
        return True
    return location.lower() in (text or "").lower()


def _job_specificity_ok(text: str, job: CanonicalJob | None) -> bool:
    if job is None:
        return True
    lowered = text.lower()
    company = _clean(getattr(job, "company", "")).lower()
    title = _clean(getattr(job, "title", "")).lower()
    title_tokens = [token for token in re.split(r"\W+", title) if len(token) > 3]
    return bool(company and company in lowered and any(token in lowered for token in title_tokens))


def _has_leadership_or_jsa_if_expected(profile: dict, text: str) -> bool:
    profile_text = _json_text(profile).lower()
    expects = any(term in profile_text for term in ("job search assistant", "leadership", "coordinated", "department head"))
    if not expects:
        return True
    lowered = text.lower()
    return any(term in lowered for term in ("job search assistant", "leadership", "coordinat", "automation", "workflow"))

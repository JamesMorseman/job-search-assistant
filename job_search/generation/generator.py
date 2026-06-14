"""Document generation — tailored resume + cover letter via provider-neutral LLM API."""

from __future__ import annotations

import json
import logging
import re
from copy import deepcopy
from pathlib import Path

import yaml
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from job_search.config import settings
from job_search.evidence import EvidencePacket, EvidenceSelector
from job_search.llm import get_llm_provider, resolve_service_config
from job_search.llm.types import LLMMessage, LLMRequest
from job_search.models import CanonicalJob

from .audit import audit_generated_documents
from .keywords import KeywordExtractor

logger = logging.getLogger(__name__)

ROOT_DIR = Path(__file__).resolve().parents[2]
RESUME_RENDERING_SPEC_PATH = ROOT_DIR / "Templates" / "resume" / "resume_rendering_spec.md"
RESUME_CONTENT_RULES_PATH = ROOT_DIR / "Templates" / "resume" / "resume_content_rules.md"
COVER_LETTER_STYLE_GUIDE_PATH = ROOT_DIR / "Templates" / "cover_letter" / "cover_letter_style_guide.md"

MAX_RESUME_SKILLS = 14
MAX_RESUME_EXPERIENCE = 2
MAX_RESUME_PROJECTS = 3
MAX_RESUME_WORK_BULLETS = 3
MAX_RESUME_PROJECT_BULLETS = 5
MAX_RESUME_COURSEWORK = 6
MAX_RESUME_COURSEWORK_EXPANDED = 9
COURSEWORK_ROW_SIZE = 3
RESUME_TARGET_WORD_FLOOR = 660
RESUME_PAGE_UTILIZATION_TARGET = 0.75
RESUME_SOFT_WORD_CAP = 760
RESUME_HARD_WORD_CAP = 950
RESUME_ONE_PAGE_WORD_CAP = 700
RESUME_ONE_PAGE_COURSEWORK_CAP = 6
RESUME_ONE_PAGE_SKILL_CAP = 14
RESUME_ONE_PAGE_SUMMARY_WORD_CAP = 60
RESUME_RIGHT_TAB_INCHES = 7.5
RESUME_CONSTRUCTION_MANAGEMENT_PROJECT_NAME = "Bridge Replacement Construction Management Planning"
RESUME_STORMWATER_PROJECT_NAME = "Stormwater Detention / Cistern Design Project"
COVER_LINE_SPACING = 1.15
COVER_HEADER_AFTER_PT = 3
COVER_CONTACT_AFTER_PT = 16
COVER_DATE_AFTER_PT = 10
COVER_EMPLOYER_AFTER_PT = 14
COVER_SALUTATION_AFTER_PT = 11
COVER_BODY_AFTER_PT = 11
COVER_CLOSING_BEFORE_PT = 14
COVER_CLOSING_AFTER_PT = 0
COVER_SIGNATURE_BEFORE_PT = 4
COVER_SIGNATURE_AFTER_PT = 8
TEMPLATE_PLACEHOLDER_RE = re.compile(r"\b[\w-]*placeholder[\w-]*\b", re.IGNORECASE)
TEMPLATE_MARKER_LABELS = (
    "company",
    "role",
    "title",
    "closing",
    "signature",
    "salutation",
    "body",
    "body_paragraph",
    "body_paragraphs",
    "hiring_manager",
    "date",
    "location",
    "candidate",
    "name",
    "email",
    "phone",
    "linkedin",
    "github",
)
TEMPLATE_MARKER_LABEL_RE = "|".join(TEMPLATE_MARKER_LABELS)
UNRESOLVED_TEMPLATE_MARKER_RE = re.compile(
    rf"""
    \{{\{{\s*[\w.-]+\s*\}}\}}
    |\$\{{\s*[\w.-]+\s*\}}
    |\{{\s*(?:{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*\}}
    |\[\s*(?:{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*\]
    |<\s*(?:placeholder|{TEMPLATE_MARKER_LABEL_RE})[\w\s.-]*>
    |(?:^|\n)\s*(?:{TEMPLATE_MARKER_LABEL_RE})\s*(?:\n|$)
    |(?:^|[\s.])[\[\{{]\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)
SCHOOL_LOCATION_FALLBACKS = {
    "farmingdale state college": "Farmingdale, New York",
    "suffolk county community college": "Selden, New York",
}


RESUME_SYSTEM = """You are an expert civil engineering resume writer.
Your task is to produce a tailored, ATS-optimized resume from the candidate's master profile.

RULES (non-negotiable):
1. Every claim must trace to a verified fact in the master profile. No fabrication.
2. Do not inflate responsibilities, overstate proficiency, or imply unearned qualifications.
3. Return ONLY valid JSON matching the schema below. No markdown, no prose outside the schema.
4. Do not include a cover letter - that is a separate call.
5. Target one page; preserve major sections before adding extra bullet depth.
6. If meaningful work experience is available, include at least one concise Work Experience
   entry with real employer, title, dates, and location.
7. Capstone/engineering project evidence remains required when available.
8. Include Job Search Assistant as automation/data/software breadth when it is selected
   and space allows; do not frame it as civil design experience.
9. Relevant Coursework is optional, compact, and capped to the strongest directly relevant
   items.

PROFILE FRAGMENT GUIDANCE:
- Treat resume_bullet_bank, role_specific_fragments, job_matching_keywords, capstone_project,
  relevant_coursework_by_category, technical_skills, software_tools, and engineering_methods
  as reusable evidence pools.
- Select only fragments and facts that match the target job description, discipline, and
  seniority level.
- Adapt selected fragments naturally for the role; do not paste them mechanically or include
  every available fact.
- Prioritize concise bullet evidence with specific projects, methods, software, coursework,
  and quantified leadership or operations details when relevant.
- If a keyword or requirement is not supported by the master profile, omit it or frame it only
  as learnability where the profile explicitly supports that.

OUTPUT SCHEMA:
{
  "professional_summary": "<2-3 sentence positioning statement for THIS role; 50-75 words max>",
  "skills": ["<skill; max 14 compact role-relevant entries>", ...],
  "education": [{"degree": "", "major": "", "institution": "", "graduation": "", "gpa": null, "honors": [], "relevant_coursework": []}],
  "experience": [{"employer": "", "title": "", "dates": "", "location": "", "bullets": ["max 1-2 concise bullets unless target role requires more"]}],
  "projects": [{"name": "", "role": "", "date": "", "bullets": ["max 3-5 concise bullets"]}],
  "activities": [{"organization": "", "role": "", "dates": "", "bullets": ["..."]}],
  "certifications": ["<cert or status>"],
  "keyword_notes": {"included": ["..."], "rationale": "<brief>"}
}"""

COVER_LETTER_SYSTEM = """You are an expert civil engineering career advisor.
Write a tailored cover letter for the candidate applying to this specific role.

RULES:
1. Every claim traces to the master profile. No fabrication.
2. The letter is human-facing — professional, readable, not keyword-stuffed.
3. Weave in ONLY top-tier JD keywords naturally; do not pad for density.
4. The letter is NOT a restatement of the resume. It positions WHY this firm, WHY this role.
5. Return ONLY valid JSON: {"salutation": "...", "body_paragraphs": ["...", "..."], "closing": "..."}
6. Use exactly 3 body paragraphs by default. Use a fourth paragraph only when role-specific context truly requires it.
7. Follow professional business-letter conventions: one page max, no bullets unless explicitly
   justified, professional tone, no exaggerated AI-style language.
8. Paragraph 1: role/company fit and concise reason for interest.
   Paragraph 2: capstone and technical engineering evidence.
   Paragraph 3: leadership/management plus Job Search Assistant evidence when relevant, framed as
   automation, workflow design, data management, testing, and technical initiative.
   Optional paragraph 4 only if necessary.

PROFILE FRAGMENT GUIDANCE:
- Treat cover_letter_fragment_bank, role_specific_fragments, job_matching_keywords,
  capstone_project, relevant_coursework_by_category, technical_skills, software_tools,
  and engineering_methods as reusable evidence pools.
- Select only role-relevant fragments based on the job description, discipline, employer
  context, and generated resume.
- Adapt fragments into a natural narrative of fit and motivation; do not paste fragments
  mechanically or repeat resume bullets.
- Use project, coursework, technical, coordination, documentation, and field-practices evidence only when it helps
  explain why the candidate fits this specific role.
- Never mention the process of using ChatGPT, Codex, or AI assistance to write the letter.
- Discuss the Job Search Assistant as evidence of skill and initiative, not as a random aside.
- Do not include unsupported claims or imply experience with tools, credentials, or duties
  not verified in the master profile."""


class DocumentGenerator:
    def __init__(self, llm_provider=None):
        self.config = resolve_service_config("generation")
        self.llm = llm_provider or get_llm_provider("generation")
        self.model = self.config.model
        self.extractor = KeywordExtractor()
        self.evidence_selector = EvidenceSelector()
        self._profile: dict | None = None

    def load_profile(self, path: str | None = None) -> None:
        profile_path = path or settings.PROFILE_PATH
        if not Path(profile_path).exists():
            raise FileNotFoundError(
                f"Profile not found at {profile_path}. "
                f"Copy the template:  cp {settings.PROFILE_TEMPLATE_PATH} {profile_path}  "
                f"and fill in James's real data before running."
            )
        with open(profile_path) as f:
            self._profile = yaml.safe_load(f)
        logger.info("Profile loaded from %s", profile_path)

    def generate(self, job: CanonicalJob) -> dict:
        """Generate tailored resume + cover letter for a job. Returns metadata dict."""
        if not self._profile:
            self.load_profile()

        jd = job.description_normalized or job.description_raw or ""
        keywords = self.extractor.extract(jd)
        profile_context, evidence_packet, used_fallback = self._build_profile_context(job, jd, keywords)

        resume_json = self._generate_resume(job, jd, keywords, profile_context)
        cover_json = self._generate_cover_letter(job, jd, keywords, resume_json, profile_context)

        resume_text = self._render_resume_text(resume_json)
        cover_text = self._render_cover_text(cover_json)
        self._reject_template_placeholders(cover_text, "cover_letter_text")

        coverage, hits, misses = self.extractor.compute_coverage(keywords, resume_text)
        audit = audit_generated_documents(
            resume_json=resume_json,
            cover_letter_json=cover_json,
            profile=self._profile or {},
            job=job,
            resume_text=resume_text,
            cover_letter_text=cover_text,
        )

        return {
            "resume_json": resume_json,
            "cover_letter_json": cover_json,
            "resume_text": resume_text,
            "cover_letter_text": cover_text,
            "document_audit": audit,
            "keyword_coverage": coverage,
            "keywords_hit": hits,
            "keywords_missed": misses,
            "evidence_packet": evidence_packet.to_prompt_dict() if evidence_packet else None,
            "evidence_fallback_used": used_fallback,
        }

    def _generate_resume(self, job: CanonicalJob, jd: str, keywords, profile_context: dict) -> dict:
        profile_json = json.dumps(profile_context, indent=2)
        kw_summary = ", ".join(k.keyword for k in keywords if k.tier == 1)
        style_guides = self._resume_style_guidance()

        user_prompt = f"""PROFILE CONTEXT:
{profile_json}

RESUME STYLE AND CONTENT RULES:
{style_guides}

TARGET ROLE:
Company: {job.company}
Title: {job.title}
Location: {job.location_city}, {job.location_state}
ATS Type: {job.ats_type.value}

JOB DESCRIPTION:
{jd[:8000]}

TIER-1 KEYWORDS (must appear in resume, both long-form and acronym):
{kw_summary}

Produce the tailored resume JSON. Emphasize the most relevant experience and projects for this discipline.
The professional summary must be rewritten specifically for this role.
Relevant Coursework is conditional: include only role-relevant coursework when space allows; otherwise omit it from the resume so it can support the cover letter narrative."""

        resp = self.llm.generate_json(LLMRequest(
            service="generation",
            model=self.model,
            max_tokens=4096,
            messages=[
                LLMMessage(role="system", content=RESUME_SYSTEM),
                LLMMessage(role="user", content=user_prompt),
            ],
        ))

        return self._qa_resume_json(json.loads(resp.content), jd)

    def _generate_cover_letter(
        self,
        job: CanonicalJob,
        jd: str,
        keywords,
        resume_json: dict,
        profile_context: dict,
    ) -> dict:
        profile_json = json.dumps(profile_context, indent=2)
        top_keywords = ", ".join(k.keyword for k in keywords if k.tier == 1)
        style_guide = self._cover_letter_style_guidance()

        user_prompt = f"""PROFILE CONTEXT:
{profile_json}

COVER LETTER STYLE GUIDE:
{style_guide}

TAILORED RESUME ALREADY GENERATED (use for consistency — do not repeat):
{json.dumps(resume_json, indent=2)[:3000]}

TARGET ROLE:
Company: {job.company}
Title: {job.title}
Location: {job.location_city}, {job.location_state}

JOB DESCRIPTION:
{jd[:6000]}

TOP KEYWORDS (weave in naturally, not mechanically):
{top_keywords}

Write the cover letter JSON. Make it compelling for a human engineering hiring manager.
Use exactly 3 body paragraphs unless a fourth is necessary:
1. role/company fit,
2. capstone/technical engineering evidence,
3. leadership/management plus Job Search Assistant evidence when relevant.
Do not mention ChatGPT, Codex, or the writing/generation process. Do not summarize the resume — make the case for fit."""

        resp = self.llm.generate_json(LLMRequest(
            service="generation",
            model=self.model,
            max_tokens=2048,
            messages=[
                LLMMessage(role="system", content=COVER_LETTER_SYSTEM),
                LLMMessage(role="user", content=user_prompt),
            ],
        ))

        return self._qa_cover_letter_json(json.loads(resp.content))

    def _build_profile_context(
        self,
        job: CanonicalJob,
        jd: str,
        keywords,
    ) -> tuple[dict, EvidencePacket | None, bool]:
        try:
            packet = self.evidence_selector.select(self._profile or {}, job, jd, keywords)
        except Exception as exc:  # noqa: BLE001 - generation must keep a safe fallback
            logger.warning("Evidence selection failed; falling back to full profile: %s", exc)
            return {"full_profile_fallback": self._profile or {}}, None, True

        if not packet.has_rich_evidence():
            logger.warning("Evidence packet too sparse; falling back to full profile")
            return {"full_profile_fallback": self._profile or {}}, packet, True

        return {
            "baseline_profile_facts": self._baseline_profile_facts(self._profile or {}),
            "selected_evidence_packet": packet.to_prompt_dict(),
            "generation_rule": (
                "Use only baseline facts and selected evidence. Preserve section/source labels for grounding. "
                "Do not infer unsupported tools, credentials, duties, or experience."
            ),
        }, packet, False

    @staticmethod
    def _baseline_profile_facts(profile: dict) -> dict:
        keys = [
            "identity",
            "eligibility",
            "relocation",
            "education",
            "education_detail",
            "certifications",
            "experience",
            "profile_summary",
            "constraints_or_todos",
        ]
        return {key: profile.get(key) for key in keys if profile.get(key) is not None}

    @classmethod
    def _resume_style_guidance(cls) -> str:
        return "\n\n".join([
            cls._read_style_doc(RESUME_RENDERING_SPEC_PATH),
            cls._read_style_doc(RESUME_CONTENT_RULES_PATH),
        ])

    @classmethod
    def _cover_letter_style_guidance(cls) -> str:
        return cls._read_style_doc(COVER_LETTER_STYLE_GUIDE_PATH)

    @staticmethod
    def _read_style_doc(path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8")
        except FileNotFoundError:
            logger.warning("Style guide not found: %s", path)
            return f"Style guide unavailable at {path}; use compact professional formatting."

    def _qa_resume_json(self, data: dict, jd: str) -> dict:
        cleaned = deepcopy(data)
        warnings: list[str] = []

        cleaned["professional_summary"] = self._clean_text(cleaned.get("professional_summary", ""))
        cleaned["skills"] = self._prioritize_resume_skills(self._clean_string_list(cleaned.get("skills", [])))[:MAX_RESUME_SKILLS]
        cleaned["education"] = self._clean_education(cleaned.get("education", []), jd, warnings)
        cleaned["experience"] = self._clean_items(
            cleaned.get("experience", []),
            required_any=("employer", "title"),
            bullet_limit=MAX_RESUME_WORK_BULLETS,
        )[:MAX_RESUME_EXPERIENCE]
        cleaned["projects"] = self._clean_items(
            cleaned.get("projects", []),
            required_any=("name", "role"),
            bullet_limit=MAX_RESUME_PROJECT_BULLETS,
        )
        cleaned["activities"] = self._clean_items(
            cleaned.get("activities", []),
            required_any=("organization", "role"),
            bullet_limit=1,
        )[:1]
        cleaned["certifications"] = self._clean_string_list(cleaned.get("certifications", []))[:3]

        self._preserve_project_identities(cleaned)
        self._allocate_resume_content(cleaned, warnings, jd)
        self._restore_profile_work_experience_if_omitted(cleaned, warnings)
        self._trim_resume_to_length(cleaned, warnings)
        self._enforce_one_page_resume(cleaned, warnings)
        qa = self._resume_qa(cleaned)
        qa["warnings"].extend(warnings)
        cleaned["rendering_qa"] = qa
        return cleaned

    def _qa_cover_letter_json(self, data: dict) -> dict:
        cleaned = deepcopy(data)
        cleaned["salutation"] = self._clean_text(cleaned.get("salutation", "Dear Hiring Manager,"))
        raw_paragraphs = cleaned.get("body_paragraphs", [])
        self._reject_template_placeholders(self._raw_cover_paragraph_text(raw_paragraphs), "cover_letter_raw_body")
        paragraphs = self._clean_cover_body_paragraphs(raw_paragraphs)
        cleaned["body_paragraphs"] = paragraphs[:4]
        cleaned["closing"] = self._clean_cover_closing(cleaned.get("closing", "Sincerely,\nJames Morseman"))
        self._validate_cover_letter_json(cleaned, raw_paragraphs)
        cleaned["rendering_diagnostics"] = {
            "raw_body_paragraph_count": len(raw_paragraphs) if isinstance(raw_paragraphs, list) else 0,
            "body_paragraph_count": len(cleaned["body_paragraphs"]),
            "multi_paragraph_output": len(cleaned["body_paragraphs"]) > 1,
            "closing_signature_separate": "\n" in cleaned["closing"],
            "template_placeholders_present": self._contains_unresolved_template_marker(
                "\n".join([cleaned["salutation"], *cleaned["body_paragraphs"], cleaned["closing"]])
            ),
        }
        return cleaned

    @staticmethod
    def _raw_cover_paragraph_text(raw_paragraphs) -> str:
        if isinstance(raw_paragraphs, str):
            return raw_paragraphs
        if isinstance(raw_paragraphs, list):
            return "\n".join(str(item) for item in raw_paragraphs)
        return str(raw_paragraphs or "")

    def _clean_cover_body_paragraphs(self, raw_paragraphs) -> list[str]:
        if isinstance(raw_paragraphs, str):
            raw_items = re.split(r"\n\s*\n+", raw_paragraphs)
        elif isinstance(raw_paragraphs, list):
            raw_items = raw_paragraphs
        else:
            raw_items = []

        paragraphs: list[str] = []
        closing_re = re.compile(r"\b(Sincerely|Regards|Best regards|Respectfully),?\b", re.IGNORECASE)
        for item in raw_items:
            text = self._clean_text(item)
            if not text:
                continue
            text = closing_re.split(text, maxsplit=1)[0].strip()
            text = re.sub(r"\bJames\s+(?:Robert\s+)?Morseman\b\s*$", "", text, flags=re.IGNORECASE).strip()
            if text:
                paragraphs.append(text)
        return paragraphs

    def _validate_cover_letter_json(self, cleaned: dict, raw_paragraphs) -> None:
        rendered = "\n".join([
            cleaned.get("salutation", ""),
            *cleaned.get("body_paragraphs", []),
            cleaned.get("closing", ""),
        ])
        self._reject_template_placeholders(rendered, "cover_letter_json")
        if not isinstance(raw_paragraphs, (list, str)):
            raise ValueError("Cover letter body_paragraphs must be a list of paragraphs or newline-separated text.")
        paragraph_count = len(cleaned.get("body_paragraphs", []))
        if paragraph_count < 3:
            raise ValueError("Cover letter must contain at least 3 separate body paragraphs.")
        if paragraph_count > 4:
            raise ValueError("Cover letter must contain no more than 4 body paragraphs.")
        if "\n" not in cleaned.get("closing", ""):
            raise ValueError("Cover letter closing and signature must be separate lines.")

    def _clean_education(self, education: list, jd: str, warnings: list[str]) -> list[dict]:
        out: list[dict] = []
        for edu in education:
            if not isinstance(edu, dict):
                continue
            item = {k: v for k, v in edu.items() if v not in ("", None, [], {})}
            if not any(item.get(k) for k in ("degree", "major", "institution")):
                continue
            coursework = self._clean_string_list(item.get("relevant_coursework", []))
            if coursework and self._coursework_is_resume_relevant(coursework, jd):
                item["relevant_coursework"] = coursework[:MAX_RESUME_COURSEWORK]
            else:
                if coursework:
                    warnings.append("Relevant coursework omitted from resume; use as cover-letter support if useful.")
                item.pop("relevant_coursework", None)
            out.append(item)
        return out[:2]

    @staticmethod
    def _coursework_is_resume_relevant(coursework: list[str], jd: str) -> bool:
        jd_norm = jd.lower()
        if not jd_norm.strip():
            return False
        role_terms = {
            "structural", "steel", "concrete", "foundation", "geotechnical", "construction",
            "transportation", "highway", "traffic", "water", "wastewater", "stormwater",
            "hydraulics", "survey", "materials", "civil", "entry level", "junior", "eit",
        }
        if not any(term in jd_norm for term in role_terms):
            return False
        course_terms = {word.lower() for course in coursework for word in course.replace("/", " ").split()}
        return bool(course_terms & set(jd_norm.replace(",", " ").replace(".", " ").split()))

    def _clean_items(self, items: list, required_any: tuple[str, ...], bullet_limit: int) -> list[dict]:
        out: list[dict] = []
        for item in items:
            if not isinstance(item, dict):
                continue
            cleaned = {k: v for k, v in item.items() if v not in ("", None, [], {})}
            if not any(cleaned.get(k) for k in required_any):
                continue
            bullets = self._clean_string_list(cleaned.get("bullets", []))[:bullet_limit]
            if bullets:
                cleaned["bullets"] = bullets
            else:
                cleaned.pop("bullets", None)
            if len(cleaned) > 0:
                out.append(cleaned)
        return out

    @staticmethod
    def _clean_string_list(values) -> list[str]:
        if not isinstance(values, list):
            return []
        return [DocumentGenerator._sanitize_resume_text(str(v).strip()) for v in values if str(v).strip()]

    @staticmethod
    def _clean_text(value) -> str:
        return DocumentGenerator._sanitize_resume_text(" ".join(str(value or "").split()))

    @staticmethod
    def _sanitize_resume_text(value: str) -> str:
        return value.replace(" ? ", "; ").replace("?", "").replace("\ufffd", "")

    def _trim_resume_to_length(self, data: dict, warnings: list[str]) -> None:
        if self._resume_word_count(data) <= RESUME_HARD_WORD_CAP:
            return

        if self._trim_coursework_row(data):
            warnings.append("Relevant coursework trimmed by full table row during length trim.")
        if self._resume_word_count(data) <= RESUME_HARD_WORD_CAP:
            return
        data["skills"] = data.get("skills", [])[:10]
        for exp in data.get("experience", []):
            exp["bullets"] = exp.get("bullets", [])[:2]
        if self._resume_word_count(data) <= RESUME_HARD_WORD_CAP:
            return
        third_project = data.get("projects", [])[2:3]
        if third_project:
            third_project[0]["bullets"] = third_project[0].get("bullets", [])[:1]
            warnings.append("Third project reduced to one bullet during length trim.")
        if self._resume_word_count(data) <= RESUME_HARD_WORD_CAP:
            return
        if len(data.get("projects", [])) > 2:
            removed = data["projects"].pop()
            warnings.append(f"Third project removed before Work Experience during length trim: {removed.get('name', '')}".strip())
        if self._resume_word_count(data) <= RESUME_HARD_WORD_CAP:
            return
        data["professional_summary"] = self._truncate_text_at_word_limit(data.get("professional_summary", ""), 60)

    def _enforce_one_page_resume(self, data: dict, warnings: list[str]) -> None:
        """Apply Phase 1 one-page policy before DOCX rendering."""
        trim_notes: list[str] = []

        def still_over() -> bool:
            return self._resume_estimated_page_count(data) > 1

        if not still_over():
            return

        if self._limit_coursework(data, RESUME_ONE_PAGE_COURSEWORK_CAP):
            trim_notes.append("excess coursework")
            data.setdefault("rendering_allocation", {}).pop("expanded_coursework", None)
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        certs = self._clean_string_list(data.get("certifications", []))
        if len(certs) > 1:
            data["certifications"] = certs[:1]
            trim_notes.append("professional development extras")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        summary_words = self._clean_text(data.get("professional_summary", "")).split()
        if len(summary_words) > RESUME_ONE_PAGE_SUMMARY_WORD_CAP:
            data["professional_summary"] = self._truncate_text_at_word_limit(
                data.get("professional_summary", ""),
                RESUME_ONE_PAGE_SUMMARY_WORD_CAP,
            )
            trim_notes.append("summary detail")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        skills = self._clean_string_list(data.get("skills", []))
        if len(skills) > RESUME_ONE_PAGE_SKILL_CAP:
            data["skills"] = self._prioritize_resume_skills(skills)[:RESUME_ONE_PAGE_SKILL_CAP]
            trim_notes.append("lower-priority skills")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        for project in reversed(data.get("projects", [])):
            if self._is_capstone_project(project) or self._is_job_search_assistant_project(project):
                continue
            bullets = project.get("bullets", [])
            if len(bullets) > 1:
                project["bullets"] = bullets[:1]
                trim_notes.append("lowest-value project bullets")
                if not still_over():
                    self._record_one_page_trim(data, warnings, trim_notes)
                    return

        for exp in data.get("experience", []):
            bullets = exp.get("bullets", [])
            if len(bullets) > 2:
                exp["bullets"] = bullets[:2]
                trim_notes.append("work-experience bullets")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        jsa = next((p for p in data.get("projects", []) if self._is_job_search_assistant_project(p)), None)
        if jsa and len(jsa.get("bullets", [])) > 1:
            jsa["bullets"] = jsa.get("bullets", [])[:1]
            trim_notes.append("Job Search Assistant extra bullets")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        capstone = next((p for p in data.get("projects", []) if self._is_capstone_project(p)), None)
        if capstone and len(capstone.get("bullets", [])) > 4:
            capstone["bullets"] = capstone.get("bullets", [])[:4]
            trim_notes.append("capstone extra bullets")
        if not still_over():
            self._record_one_page_trim(data, warnings, trim_notes)
            return

        if len(data.get("projects", [])) > 2:
            protected = [p for p in data.get("projects", []) if self._is_capstone_project(p) or self._is_job_search_assistant_project(p)]
            unprotected = [p for p in data.get("projects", []) if p not in protected]
            if unprotected:
                data["projects"] = protected + unprotected[:1]
                trim_notes.append("extra project entries")

        self._record_one_page_trim(data, warnings, trim_notes)

    def _record_one_page_trim(self, data: dict, warnings: list[str], trim_notes: list[str]) -> None:
        data.setdefault("rendering_allocation", {})["one_page_enforced"] = True
        if trim_notes:
            warnings.append("One-page resume enforcement trimmed: " + ", ".join(dict.fromkeys(trim_notes)) + ".")

    def _truncate_text_at_word_limit(self, text: str, word_limit: int) -> str:
        cleaned = self._clean_text(text)
        words = cleaned.split()
        if len(words) <= word_limit:
            return cleaned
        clipped = " ".join(words[:word_limit]).strip()
        sentence_match = re.match(r"^(.+[.!?])(?:\s+[^.!?]*)?$", clipped)
        if sentence_match:
            return sentence_match.group(1).strip()
        return clipped

    def _limit_coursework(self, data: dict, limit: int) -> bool:
        changed = False
        remaining = max(limit, 0)
        for edu in data.get("education", []):
            coursework = self._clean_string_list(edu.get("relevant_coursework", []))
            if not coursework:
                continue
            keep = coursework[:remaining]
            remaining = max(0, remaining - len(keep))
            if len(keep) < len(coursework):
                changed = True
            if keep:
                edu["relevant_coursework"] = keep
            else:
                edu.pop("relevant_coursework", None)
        return changed

    def _prioritize_resume_skills(self, skills: list[str]) -> list[str]:
        compacted = self._clean_string_list(skills)
        seen: set[str] = set()
        deduped: list[str] = []
        for skill in compacted:
            key = skill.lower()
            if key in seen:
                continue
            deduped.append(skill)
            seen.add(key)
        return sorted(deduped, key=lambda skill: (self._resume_skill_priority(skill), deduped.index(skill)))

    @staticmethod
    def _resume_skill_priority(skill: str) -> int:
        normalized = skill.lower()
        high_value_terms = (
            "ram structural system",
            "autocad",
            "revit",
            "excel",
            "google sheets",
            "sqlite",
            "python",
            "automation",
            "microsoft project",
            "project scheduling",
        )
        if any(term in normalized for term in high_value_terms):
            return 0
        standards = ("asce", "aisc", "aci", "astm", "osha")
        if any(term in normalized for term in standards):
            return 1
        engineering_terms = (
            "structural",
            "steel",
            "concrete",
            "foundation",
            "construction",
            "stormwater",
            "hydraulic",
            "civil",
        )
        if any(term in normalized for term in engineering_terms):
            return 2
        return 3

    def _trim_coursework_row(self, data: dict) -> bool:
        coursework: list[str] = []
        for edu in data.get("education", []):
            coursework.extend(self._clean_string_list(edu.get("relevant_coursework", [])))
        if not coursework:
            return False
        rows = (len(coursework) + COURSEWORK_ROW_SIZE - 1) // COURSEWORK_ROW_SIZE
        keep = max(0, (rows - 1) * COURSEWORK_ROW_SIZE)
        kept_items = set(coursework[:keep])
        for edu in data.get("education", []):
            current = self._clean_string_list(edu.get("relevant_coursework", []))
            if not current:
                continue
            remaining = [course for course in current if course in kept_items]
            if remaining:
                edu["relevant_coursework"] = remaining
            else:
                edu.pop("relevant_coursework", None)
        return True

    def _allocate_resume_content(self, data: dict, warnings: list[str], jd: str) -> None:
        original_experience = deepcopy(data.get("experience", []))
        original_projects = deepcopy(data.get("projects", []))
        data["projects"] = self._allocate_projects(deepcopy(original_projects), warnings)
        if data.get("experience"):
            data["experience"] = data["experience"][:1]
            data["experience"][0]["bullets"] = data["experience"][0].get("bullets", [])[:2]
        for index, project in enumerate(data.get("projects", [])):
            if self._is_capstone_project(project):
                project["bullets"] = project.get("bullets", [])[:3]
            else:
                project["bullets"] = project.get("bullets", [])[:1 if data.get("experience") else 2]
        self._expand_underfilled_resume(data, original_experience, original_projects, warnings, jd)

    def _restore_profile_work_experience_if_omitted(self, data: dict, warnings: list[str]) -> None:
        if data.get("experience") or not self._profile:
            return
        if self._resume_word_count(data) >= RESUME_SOFT_WORD_CAP:
            return
        profile_experience = self._profile_work_experience_items()
        if not profile_experience:
            return
        data["experience"] = profile_experience[:1]
        warnings.append("Work Experience restored from baseline profile facts after generated resume omitted it.")

    def _profile_work_experience_items(self) -> list[dict]:
        if not self._profile:
            return []
        out: list[dict] = []
        for entry in self._profile.get("experience", []) or []:
            if not isinstance(entry, dict):
                continue
            employer = self._clean_text(entry.get("employer", ""))
            title = self._clean_text(entry.get("title", ""))
            if not employer and not title:
                continue
            item = {
                "employer": employer,
                "title": title,
                "dates": self._format_profile_date_range(entry),
                "location": self._clean_text(str(entry.get("location", "")).replace(";", " | ")),
                "bullets": self._profile_work_bullets(entry)[:MAX_RESUME_WORK_BULLETS],
            }
            out.append({k: v for k, v in item.items() if v not in ("", [], None)})
        return out

    def _profile_work_bullets(self, entry: dict) -> list[str]:
        bullets: list[str] = []
        for bullet in entry.get("bullets", []) or []:
            if isinstance(bullet, dict):
                text = bullet.get("text", "")
            else:
                text = bullet
            clean = self._clean_text(text)
            if clean:
                bullets.append(clean)
        return bullets

    def _profile_project_items(self) -> list[dict]:
        if not self._profile:
            return []
        out: list[dict] = []
        for entry in self._profile.get("projects", []) or []:
            if not isinstance(entry, dict):
                continue
            name = self._clean_text(entry.get("name", ""))
            role = self._clean_text(entry.get("role", ""))
            if not name:
                continue
            bullets = self._profile_project_bullets(entry)
            item = {
                "name": name,
                "role": role or self._clean_text(str(entry.get("type", "")).replace("_", " ").title()),
                "date": self._format_year_month(entry.get("date")),
                "bullets": bullets[:2],
            }
            out.append({k: v for k, v in item.items() if v not in ("", [], None)})
        return out

    def _profile_project_bullets(self, entry: dict) -> list[str]:
        bullets = self._profile_work_bullets(entry)
        if len(bullets) < 2:
            return bullets
        first = bullets[0].rstrip(".")
        second = bullets[1][0].lower() + bullets[1][1:] if bullets[1] else bullets[1]
        return [f"{first}; {second}", *bullets[2:]]

    def _format_profile_date_range(self, entry: dict) -> str:
        start = self._format_year_month(entry.get("start_date"))
        end_raw = entry.get("end_date")
        end = "Present" if str(end_raw).lower() == "present" else self._format_year_month(end_raw)
        return " - ".join(part for part in [start, end] if part)

    @staticmethod
    def _format_year_month(value) -> str:
        text = str(value or "").strip()
        if len(text) != 7 or text[4] != "-":
            return text
        months = {
            "01": "January", "02": "February", "03": "March", "04": "April",
            "05": "May", "06": "June", "07": "July", "08": "August",
            "09": "September", "10": "October", "11": "November", "12": "December",
        }
        year, month = text.split("-")
        return f"{months.get(month, month)} {year}"

    def _expand_underfilled_resume(
        self,
        data: dict,
        original_experience: list[dict],
        original_projects: list[dict],
        warnings: list[str],
        jd: str,
    ) -> None:
        if self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        expansion_order: list[str] = []

        def expand(target: dict, source: dict | None, limit: int, label: str) -> None:
            if not source or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
                return
            before = len(target.get("bullets", []))
            self._expand_bullets(target, source.get("bullets", []), limit)
            if len(target.get("bullets", [])) > before:
                expansion_order.append(label)

        capstone = next((p for p in data.get("projects", []) if self._is_capstone_project(p)), None)
        expand(capstone, self._matching_project(original_projects, capstone), MAX_RESUME_PROJECT_BULLETS, "capstone_bullets")

        self._expand_profile_bullet_bank(data, expansion_order)

        jsa = next((p for p in data.get("projects", []) if self._is_job_search_assistant_project(p)), None)
        expand(jsa, self._matching_project(original_projects, jsa), 3, "job_search_assistant_bullets")

        for project in data.get("projects", []):
            if self._is_capstone_project(project) or self._is_job_search_assistant_project(project):
                continue
            expand(project, self._matching_project(original_projects, project), 2, "academic_project_bullets")

        self._restore_profile_relevant_project(data, jd, expansion_order)
        self._restore_profile_breadth_project(data, expansion_order)

        if data.get("experience"):
            current = data["experience"][0]
            source = self._matching_item(original_experience, current, ("employer", "title"))
            expand(current, source, MAX_RESUME_WORK_BULLETS, "work_experience_bullets")

        self._expand_profile_work_experience(data, expansion_order)
        self._expand_profile_coursework(data, expansion_order)
        self._expand_profile_skills(data, expansion_order)
        self._expand_profile_certifications(data, expansion_order)
        self._expand_profile_summary(data, expansion_order)

        if expansion_order:
            data.setdefault("rendering_allocation", {})["expansion_order"] = expansion_order
        if self._resume_word_count(data) < RESUME_TARGET_WORD_FLOOR:
            warnings.append("Resume remains below page-utilization target after safe expansion.")

    def _expand_profile_bullet_bank(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        bank = self._profile.get("resume_bullet_bank", {}) or {}
        if not isinstance(bank, dict):
            return
        capstone = next((p for p in data.get("projects", []) if self._is_capstone_project(p)), None)
        if capstone:
            before = len(capstone.get("bullets", []))
            self._expand_bullets(capstone, bank.get("structural", []) + bank.get("leadership", []), MAX_RESUME_PROJECT_BULLETS)
            if len(capstone.get("bullets", [])) > before and "capstone_bullets" not in expansion_order:
                expansion_order.append("capstone_bullets")

        if self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        jsa = next((p for p in data.get("projects", []) if self._is_job_search_assistant_project(p)), None)
        if jsa:
            before = len(jsa.get("bullets", []))
            self._expand_bullets(jsa, bank.get("software_data_automation", []), 3)
            if len(jsa.get("bullets", [])) > before and "job_search_assistant_bullets" not in expansion_order:
                expansion_order.append("job_search_assistant_bullets")

    def _restore_profile_breadth_project(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        projects = data.setdefault("projects", [])
        if len(projects) >= MAX_RESUME_PROJECTS:
            return
        existing = next((project for project in projects if self._is_construction_management_project(project)), None)
        if existing:
            self._canonicalize_construction_management_project(existing)
            return
        construction_project = next(
            (project for project in self._profile_project_items() if self._is_construction_management_project(project)),
            None,
        )
        if construction_project:
            self._canonicalize_construction_management_project(construction_project)
            projects.append(construction_project)
            expansion_order.append("academic_project_bullets")

    def _restore_profile_relevant_project(self, data: dict, jd: str, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        projects = data.setdefault("projects", [])
        if len(projects) >= MAX_RESUME_PROJECTS:
            return
        existing_topics = {self._project_topic(project) for project in projects}
        for project in self._profile_project_items():
            if self._is_capstone_project(project) or self._is_job_search_assistant_project(project):
                continue
            topic = self._project_topic(project)
            if topic in existing_topics:
                continue
            if not self._profile_project_matches_job(project, jd):
                continue
            if self._is_construction_management_project(project):
                self._canonicalize_construction_management_project(project)
            if self._is_stormwater_project(project):
                self._canonicalize_stormwater_project(project)
            projects.append(project)
            expansion_order.append("academic_project_bullets")
            return

    def _preserve_project_identities(self, data: dict) -> None:
        for project in data.get("projects", []):
            if self._is_construction_management_project(project):
                self._canonicalize_construction_management_project(project)
            if self._is_stormwater_project(project):
                self._canonicalize_stormwater_project(project)

    def _canonicalize_construction_management_project(self, project: dict) -> None:
        project["name"] = RESUME_CONSTRUCTION_MANAGEMENT_PROJECT_NAME
        role = self._clean_text(project.get("role", ""))
        if not role or "coursework" in role.lower():
            project["role"] = "Construction Planning Project"

    def _canonicalize_stormwater_project(self, project: dict) -> None:
        project["name"] = RESUME_STORMWATER_PROJECT_NAME
        role = self._clean_text(project.get("role", ""))
        if not role or "coursework" in role.lower():
            project["role"] = "Academic Hydrology / Stormwater Project"

    def _expand_profile_work_experience(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR or not data.get("experience"):
            return
        current = data["experience"][0]
        source = self._matching_item(self._profile_work_experience_items(), current, ("employer", "title"))
        if not source:
            return
        before = len(current.get("bullets", []))
        self._expand_bullets(current, source.get("bullets", []), MAX_RESUME_WORK_BULLETS)
        if len(current.get("bullets", [])) > before and "work_experience_bullets" not in expansion_order:
            expansion_order.append("work_experience_bullets")

    def _expand_profile_coursework(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        changed = False
        for edu in data.get("education", []):
            entry = self._profile_education_entry(str(edu.get("institution", "")))
            source = self._clean_string_list((entry or {}).get("relevant_coursework", []))
            if not source:
                continue
            current = self._clean_string_list(edu.get("relevant_coursework", []))
            for course in source:
                if len(current) >= MAX_RESUME_COURSEWORK_EXPANDED or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
                    break
                if course not in current:
                    current.append(course)
                    changed = True
            if current:
                edu["relevant_coursework"] = current
        if changed:
            data.setdefault("rendering_allocation", {})["expanded_coursework"] = True
            expansion_order.append("coursework_items")

    def _expand_profile_skills(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        candidates: list[str] = []
        technical = self._profile.get("technical_skills", {}) or {}
        if not isinstance(technical, dict):
            technical = {}
        for key in ("software_data_automation", "leadership_operations", "structural", "construction", "geotechnical_foundation", "water_site_civil"):
            candidates.extend(self._clean_string_list(technical.get(key, [])))
        software_tools = self._profile.get("software_tools", {}) or {}
        tools = software_tools.get("verified", []) if isinstance(software_tools, dict) else software_tools
        candidates = self._clean_string_list(tools) + candidates

        skills = self._prioritize_resume_skills(self._clean_string_list(data.get("skills", [])))
        before = len(skills)
        seen = {skill.lower() for skill in skills}
        for skill in self._prioritize_resume_skills(candidates):
            if len(skills) >= MAX_RESUME_SKILLS or self._resume_word_count({**data, "skills": skills}) >= RESUME_TARGET_WORD_FLOOR:
                break
            key = skill.lower()
            if key not in seen:
                skills.append(skill)
                seen.add(key)
        if len(skills) > before:
            data["skills"] = self._prioritize_resume_skills(skills)
            expansion_order.append("skills_items")

    def _expand_profile_certifications(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        certs = self._clean_string_list(data.get("certifications", []))
        profile_certs = self._profile.get("certifications", {}) or {}
        if not isinstance(profile_certs, dict):
            return
        candidates = self._clean_string_list(profile_certs.get("professional_development", []))
        before = len(certs)
        seen = {cert.lower() for cert in certs}
        for cert in candidates:
            if len(certs) >= 3 or self._resume_word_count({**data, "certifications": certs}) >= RESUME_TARGET_WORD_FLOOR:
                break
            key = cert.lower()
            if key not in seen:
                certs.append(cert)
                seen.add(key)
        if len(certs) > before:
            data["certifications"] = certs
            expansion_order.append("professional_development_items")

    def _expand_profile_summary(self, data: dict, expansion_order: list[str]) -> None:
        if not self._profile or self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
            return
        profile_summary = self._profile.get("profile_summary", {}) or {}
        if not isinstance(profile_summary, dict):
            return
        current = self._clean_text(data.get("professional_summary", ""))
        candidates = [
            profile_summary.get("current_status", ""),
            profile_summary.get("positioning", ""),
            *self._clean_string_list(profile_summary.get("primary_differentiators", [])),
        ]
        sentences = [current] if current else []
        current_lower = current.lower()
        for candidate in self._clean_string_list(candidates):
            if self._resume_word_count(data) >= RESUME_TARGET_WORD_FLOOR:
                break
            if candidate.lower() in current_lower:
                continue
            sentences.append(candidate)
            summary_words = " ".join(sentences).split()
            if len(summary_words) >= 75:
                sentences = [self._truncate_text_at_word_limit(" ".join(sentences), 75)]
                break
        expanded = " ".join(sentences).strip()
        if expanded and expanded != current:
            data["professional_summary"] = expanded
            expansion_order.append("summary_detail")

    @staticmethod
    def _expand_bullets(target: dict, source_bullets, limit: int) -> None:
        if target is None:
            return
        bullets = target.setdefault("bullets", [])
        for bullet in source_bullets or []:
            if len(bullets) >= limit:
                break
            if bullet not in bullets:
                bullets.append(bullet)

    @staticmethod
    def _matching_item(items: list[dict], target: dict | None, keys: tuple[str, ...]) -> dict | None:
        if not target:
            return None
        for item in items:
            if all(item.get(key) == target.get(key) for key in keys):
                return item
        return None

    def _matching_project(self, projects: list[dict], target: dict | None) -> dict | None:
        if not target:
            return None
        return self._matching_item(projects, target, ("name",)) or self._matching_item(projects, target, ("role",))

    def _allocate_projects(self, projects: list[dict], warnings: list[str]) -> list[dict]:
        if not projects:
            return []
        capstone = next((p for p in projects if self._is_capstone_project(p)), None)
        jsa = next((p for p in projects if self._is_job_search_assistant_project(p)), None)
        selected: list[dict] = []
        if capstone:
            selected.append(capstone)
        if jsa and jsa not in selected:
            selected.append(jsa)
        for project in projects:
            if project is capstone or project is jsa:
                continue
            if capstone and self._project_topic(project) == self._project_topic(capstone):
                warnings.append(f"Redundant project omitted: {project.get('name', '')}".strip())
                continue
            selected.append(project)
            break
        return selected[:MAX_RESUME_PROJECTS]

    @staticmethod
    def _is_capstone_project(project: dict) -> bool:
        text = " ".join(str(project.get(k, "")) for k in ("name", "role", "date")).lower()
        return "capstone" in text or "baldwin" in text

    @staticmethod
    def _is_job_search_assistant_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return "job search assistant" in text or "job-search automation" in text

    @staticmethod
    def _is_construction_management_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return (
            "construction management" in text
            or "construction planning" in text
            or ("traffic control" in text and "schedule" in text)
            or ("bridge replacement" in text and "planning" in text)
        )

    @staticmethod
    def _is_stormwater_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return (
            "stormwater" in text
            or "cistern" in text
            or ("runoff" in text and "drainage" in text)
        )

    @staticmethod
    def _profile_project_matches_job(project: dict, jd: str) -> bool:
        project_text = json.dumps(project).lower()
        jd_text = str(jd or "").lower()
        matches = (
            (("transportation" in jd_text or "bridge" in jd_text or "roadway" in jd_text) and any(
                term in project_text for term in ("transportation", "bridge", "traffic control", "roadway")
            ))
            or (("construction" in jd_text or "field" in jd_text or "schedule" in jd_text) and any(
                term in project_text for term in ("construction", "schedule", "traffic control", "safety")
            ))
            or (("land development" in jd_text or "site civil" in jd_text or "stormwater" in jd_text or "drainage" in jd_text) and any(
                term in project_text for term in ("stormwater", "drainage", "site_civil", "water_resources", "cistern")
            ))
            or (("water" in jd_text or "wastewater" in jd_text) and any(
                term in project_text for term in ("water_resources", "stormwater", "cistern", "runoff")
            ))
        )
        return bool(matches)

    @staticmethod
    def _project_topic(project: dict) -> str:
        text = json.dumps(project).lower()
        if any(term in text for term in ("python", "sqlite", "openai", "llm", "automation", "data ingestion")):
            return "software_automation"
        if any(term in text for term in ("structural", "steel", "concrete", "ram", "load", "framing")):
            return "structural"
        if any(term in text for term in ("construction", "traffic", "safety", "schedule")):
            return "construction"
        if any(term in text for term in ("stormwater", "water", "hydraulic", "cistern")):
            return "water"
        return text[:40]

    @staticmethod
    def _is_software_data_heavy_resume(resume_json: dict) -> bool:
        text = " ".join([
            str(resume_json.get("professional_summary", "")),
            json.dumps(resume_json.get("keyword_notes", {})),
        ]).lower()
        civil_signals = [
            "civil engineer", "structural", "transportation", "water resources",
            "construction", "site civil", "land development", "geotechnical",
        ]
        if any(signal in text for signal in civil_signals):
            return False
        signals = [
            "data analyst", "automation analyst", "software engineer", "software developer",
            "developer role", "python automation role", "data automation role",
            "database role", "workflow automation role",
        ]
        return sum(1 for signal in signals if signal in text) >= 2

    def _resume_qa(self, data: dict) -> dict:
        word_count = self._resume_word_count(data)
        return {
            "estimated_words": word_count,
            "page_utilization_estimate": self._resume_page_utilization(data),
            "estimated_page_count": self._resume_estimated_page_count(data),
            "length_risk": (
                "hard_cap_risk" if word_count > RESUME_HARD_WORD_CAP
                else "soft_cap_risk" if word_count > RESUME_SOFT_WORD_CAP
                else "ok"
            ),
            "blank_sections_removed": True,
            "expansion_order": data.get("rendering_allocation", {}).get("expansion_order", []),
            "warnings": [],
        }

    @staticmethod
    def _resume_word_count(data: dict) -> int:
        text = json.dumps({
            k: v for k, v in data.items()
            if k not in {"keyword_notes", "rendering_qa", "rendering_allocation"}
        })
        return len(text.split())

    def _render_resume_text(self, data: dict) -> str:
        """Plain-text render of the resume JSON for keyword coverage computation."""
        parts = [data.get("professional_summary", "")]
        parts += data.get("skills", [])
        for exp in data.get("experience", []):
            parts.append(exp.get("title", ""))
            parts += exp.get("bullets", [])
        for proj in data.get("projects", []):
            parts.append(proj.get("name", ""))
            parts += proj.get("bullets", [])
        for edu in data.get("education", []):
            parts += edu.get("relevant_coursework", [])
        parts += data.get("certifications", [])
        return " ".join(parts)

    def _render_cover_text(self, data: dict) -> str:
        data = self._qa_cover_letter_json(data)
        parts = [self._clean_text(data.get("salutation", ""))]
        parts += self._clean_string_list(data.get("body_paragraphs", []))
        closing = self._clean_cover_closing(data.get("closing", ""))
        if closing:
            parts.append(closing)
        rendered = "\n\n".join(part for part in parts if part)
        self._reject_template_placeholders(rendered, "cover_letter_text")
        return rendered

    def prepare_resume_for_rendering(self, resume_json: dict) -> tuple[dict, list[str]]:
        rendered = deepcopy(resume_json)
        warnings: list[str] = []
        self._enforce_one_page_resume(rendered, warnings)
        return rendered, warnings

    def save_cover_docx(
        self,
        cover_json: dict,
        output_path: str,
        job: CanonicalJob | None = None,
        today: str | None = None,
        include_location: bool = True,
        include_linkedin: bool = False,
    ) -> None:
        """Save a cover letter as DOCX while preserving paragraph and signature structure."""
        cover_json = self._qa_cover_letter_json(cover_json)
        doc = DocxDocument()
        self._configure_cover_document(doc)
        profile = self._profile or {}
        identity = profile.get("identity", {})

        self.add_cover_name_header(doc, identity)
        self.add_cover_contact_block(
            doc,
            identity,
            include_location=include_location,
            include_linkedin=include_linkedin,
        )

        if today:
            self._add_cover_paragraph(doc, today, after=COVER_DATE_AFTER_PT)
        if job:
            employer_lines = [job.company]
            location = ", ".join(part for part in [job.location_city, job.location_state] if part)
            if location:
                employer_lines.append(location)
            for index, line in enumerate(employer_lines):
                self._add_cover_paragraph(
                    doc,
                    line,
                    after=COVER_EMPLOYER_AFTER_PT if index == len(employer_lines) - 1 else 0,
                )

        salutation = self._clean_text(cover_json.get("salutation", "Dear Hiring Manager,"))
        if salutation:
            self._add_cover_paragraph(doc, salutation, after=COVER_SALUTATION_AFTER_PT)

        for paragraph in self._clean_string_list(cover_json.get("body_paragraphs", [])):
            self._add_cover_paragraph(doc, paragraph, after=COVER_BODY_AFTER_PT)

        closing = self._clean_cover_closing(cover_json.get("closing", "Sincerely,\nJames Morseman"))
        for index, line in enumerate(closing.splitlines()):
            if index == 0:
                self._add_cover_paragraph(
                    doc,
                    line,
                    before=COVER_CLOSING_BEFORE_PT,
                    after=COVER_CLOSING_AFTER_PT,
                )
            else:
                self._add_cover_paragraph(doc, line, before=COVER_SIGNATURE_BEFORE_PT, after=COVER_SIGNATURE_AFTER_PT)

        rendered = "\n\n".join(p.text for p in doc.paragraphs if p.text)
        self._reject_template_placeholders(rendered, "cover_docx")
        doc.save(output_path)

    def save_docx(self, resume_json: dict, output_path: str) -> None:
        """Save the resume as a compact single-column .docx following the template spec."""
        resume_json, render_warnings = self.prepare_resume_for_rendering(resume_json)
        doc = DocxDocument()
        self._configure_resume_document(doc)
        profile = self._profile or {}
        identity = profile.get("identity", {})

        self.add_name_header(doc, identity)
        self.add_resume_contact_block(doc, identity)

        summary = resume_json.get("professional_summary")
        if summary:
            self.add_section_heading(doc, "Professional Summary")
            self._add_body_paragraph(doc, summary)

        education = resume_json.get("education", [])
        if education:
            self.add_section_heading(doc, "Education")
            for edu in education:
                self.add_school_entry(doc, edu)

        coursework = self._selected_coursework(resume_json)
        if coursework:
            self.add_coursework_section(doc, coursework)

        projects = resume_json.get("projects", [])
        if projects:
            self.add_section_heading(doc, "Engineering Experience")
            for project in projects:
                self.add_project_entry(doc, project)

        experience = resume_json.get("experience", [])
        if experience:
            self.add_section_heading(doc, "Work Experience")
            for entry in experience:
                self.add_work_entry(doc, entry)

        skill_groups = self._group_skills(
            resume_json.get("skills", []),
            separate_programming_data=self._is_software_data_heavy_resume(resume_json),
            resume_json=resume_json,
        )
        if skill_groups:
            self.add_section_heading(doc, "Technical Skills")
            for label, values in skill_groups.items():
                self.add_skill_group(doc, label, values)

        certs = resume_json.get("certifications", [])
        if certs:
            self.add_professional_development_line(doc, certs)

        qa = self.resume_renderer_qa(resume_json)
        qa["warnings"].extend(render_warnings)
        if qa["warnings"]:
            logger.warning("Resume renderer QA warnings: %s", qa["warnings"])

        doc.save(output_path)
        logger.info("Resume saved to %s", output_path)

    def _configure_resume_document(self, doc) -> None:
        section = doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)

        if "Resume Bullet" not in [style.name for style in doc.styles]:
            bullet = doc.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
            bullet.base_style = doc.styles["List Bullet"]
            bullet.font.name = "Calibri"
            bullet.font.size = Pt(10)
            bullet.paragraph_format.left_indent = Inches(0.5)
            bullet.paragraph_format.first_line_indent = Inches(-0.25)
            bullet.paragraph_format.line_spacing = 1.0
            bullet.paragraph_format.space_before = Pt(0)
            bullet.paragraph_format.space_after = Pt(0)

    def _configure_cover_document(self, doc) -> None:
        section = doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.line_spacing = COVER_LINE_SPACING
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)

    def add_name_header(self, doc, identity: dict) -> None:
        name = f"{identity.get('first_name', '')} {identity.get('last_name', '')}".strip()
        if not name:
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_spacing(p, line=1.0667)
        run = p.add_run(name)
        self._format_run(run, size=18)

    def add_cover_name_header(self, doc, identity: dict) -> None:
        name = f"{identity.get('first_name', '')} {identity.get('last_name', '')}".strip()
        if not name:
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_spacing(p, after=COVER_HEADER_AFTER_PT, line=COVER_LINE_SPACING)
        run = p.add_run(name)
        self._format_run(run, size=16)

    def add_contact_line(self, doc, identity: dict, include_linkedin: bool = False, include_location: bool = False) -> None:
        contact_parts = [identity.get("email"), identity.get("phone")]
        if include_linkedin:
            contact_parts.append(identity.get("linkedin_url"))
        if include_location:
            location = identity.get("location", {})
            loc_text = ", ".join([x for x in [location.get("city"), location.get("state")] if x])
            contact_parts.append(loc_text)
        contact_parts = [part for part in contact_parts if part]
        if not contact_parts:
            return
        p = doc.add_paragraph(" | ".join(contact_parts))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_spacing(p, line=1.0667)
        self._format_paragraph_runs(p, size=10)

    def add_resume_contact_block(self, doc, identity: dict) -> None:
        """Render approved resume header contact lines from profile identity."""
        primary_parts = [
            self._clean_text(identity.get("phone", "")),
            self._clean_text(identity.get("email", "")),
        ]
        self._add_centered_contact_line(doc, primary_parts)

        link_parts = [
            self._clean_text(identity.get("linkedin_url", "")),
            self._clean_text(identity.get("github_url", "")),
        ]
        self._add_centered_contact_line(doc, link_parts)

    def add_cover_contact_block(
        self,
        doc,
        identity: dict,
        include_location: bool = True,
        include_linkedin: bool = False,
    ) -> None:
        """Render stable business-letter contact line from profile identity."""
        contact_parts = [
            self._clean_text(identity.get("phone", "")),
            self._clean_text(identity.get("email", "")),
        ]
        if include_linkedin:
            contact_parts.append(self._clean_text(identity.get("linkedin_url", "")))
        if include_location:
            location = identity.get("location", {})
            loc_text = ", ".join([
                self._clean_text(location.get("city", "")),
                self._clean_text(location.get("state", "")),
            ]).strip(", ")
            contact_parts.append(loc_text)
        self._add_centered_contact_line(doc, contact_parts, after=COVER_CONTACT_AFTER_PT, line=COVER_LINE_SPACING)

    def _add_centered_contact_line(
        self,
        doc,
        parts: list[str],
        after: float = 0,
        line: float = 1.0667,
    ) -> None:
        parts = [part for part in parts if part]
        if not parts:
            return
        p = doc.add_paragraph(" | ".join(parts))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_spacing(p, after=after, line=line)
        self._format_paragraph_runs(p, size=10)

    def add_section_heading(self, doc, title: str) -> None:
        if not title:
            return
        p = doc.add_paragraph()
        self._set_spacing(p, line=1.0667)
        self._add_bottom_border(p)
        run = p.add_run(title)
        self._format_run(run, size=12, bold=True)

    def add_school_entry(self, doc, edu: dict) -> None:
        institution = self._clean_text(edu.get("institution", ""))
        if not institution:
            return
        location = self._education_location(edu)
        graduation = self._education_graduation_date(edu.get("graduation", ""))
        honors = self._education_honors(edu)

        p = doc.add_paragraph()
        self._set_spacing(p, line=1.0667)
        self._add_right_tab(p)
        self._format_run(p.add_run(institution), bold=True)
        if location:
            p.add_run(f" | {location}")
        if graduation:
            self._format_run(p.add_run(f"\t{graduation}"), italic=True)

        degree = self._degree_line(edu)
        if degree:
            p2 = doc.add_paragraph()
            self._set_spacing(p2, line=1.0667)
            p2.paragraph_format.left_indent = Inches(0.25)
            self._add_right_tab(p2)
            p2.add_run(degree)
            if honors:
                self._format_run(p2.add_run(f"\t{'; '.join(honors)}"), italic=True)
            self._format_paragraph_runs(p2, size=10)

    def add_coursework_section(self, doc, coursework: list[str]) -> None:
        if not coursework:
            return
        self.add_section_heading(doc, "Relevant Coursework")
        coursework = self._arrange_coursework_for_table(coursework)[:MAX_RESUME_COURSEWORK_EXPANDED]
        table = doc.add_table(rows=0, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        self._remove_table_borders(table)
        for idx in range(0, len(coursework), 3):
            row = table.add_row()
            for col_idx, course in enumerate(coursework[idx:idx + 3]):
                cell = row.cells[col_idx]
                paragraph = cell.paragraphs[0]
                self._set_spacing(paragraph)
                run = paragraph.add_run(course)
                self._format_run(run, size=10)

    def add_project_entry(self, doc, project: dict) -> None:
        name = self._clean_text(project.get("name", ""))
        if not name:
            return
        role = self._project_role(project)
        date = self._project_date(project)
        organization = self._project_organization(project)
        p = doc.add_paragraph()
        self._set_spacing(p, line=1.0667)
        self._add_right_tab(p)
        self._format_run(p.add_run(name), bold=True)
        if date:
            self._format_run(p.add_run(f"\t{date}"), italic=True)
        self._format_paragraph_runs(p, size=10)
        if role or organization:
            p2 = doc.add_paragraph()
            self._set_spacing(p2, line=1.0667)
            descriptor = " | ".join(part for part in [role, organization] if part)
            p2.add_run(descriptor)
            self._format_paragraph_runs(p2, size=10)
        for bullet in self._clean_string_list(project.get("bullets", [])):
            self._add_bullet(doc, bullet)

    def add_work_entry(self, doc, entry: dict) -> None:
        title = self._clean_text(entry.get("title", ""))
        employer = self._clean_text(entry.get("employer", ""))
        if not title and not employer:
            return
        dates = self._clean_text(entry.get("dates", ""))
        p = doc.add_paragraph()
        self._set_spacing(p, line=1.0667)
        self._add_right_tab(p)
        if title:
            self._format_run(p.add_run(title), bold=True)
        if employer:
            p.add_run(f" | {employer}" if title else employer)
        if dates:
            self._format_run(p.add_run(f"\t{dates}"), italic=True)
        self._format_paragraph_runs(p, size=10)

        location = self._clean_text(entry.get("location", ""))
        if location:
            loc = doc.add_paragraph()
            self._set_spacing(loc, line=1.0667)
            self._add_right_tab(loc)
            self._format_run(loc.add_run(location), italic=True)

        for bullet in self._clean_string_list(entry.get("bullets", [])):
            self._add_bullet(doc, bullet)

    def add_skill_group(self, doc, label: str, values: list[str]) -> None:
        values = self._clean_string_list(values)
        if not values:
            return
        p = doc.add_paragraph()
        self._set_spacing(p, after=0, line=1.0)
        self._add_right_tab(p)
        self._format_run(p.add_run(f"{label}:"), bold=True)
        p.add_run(" " + "; ".join(values))
        self._format_paragraph_runs(p, size=10)

    def add_professional_development_line(self, doc, certs: list[str]) -> None:
        certs = self._clean_string_list(certs)
        if not certs:
            return
        p = doc.add_paragraph()
        self._set_spacing(p, before=0, after=0, line=1.0)
        self._format_run(p.add_run("Professional Development:"), bold=True)
        p.add_run(" " + " | ".join(certs))
        self._format_paragraph_runs(p, size=10)

    def resume_renderer_qa(self, resume_json: dict) -> dict:
        warnings: list[str] = []
        sections = self._planned_resume_sections(resume_json)
        if any(not section for section in sections):
            warnings.append("blank_section")
        if "Relevant Coursework" in sections and "Education" in sections:
            if sections.index("Relevant Coursework") < sections.index("Education"):
                warnings.append("coursework_before_education")
        if "Relevant Coursework" in sections and "Engineering Experience" in sections:
            if sections.index("Relevant Coursework") > sections.index("Engineering Experience"):
                warnings.append("coursework_after_engineering_experience")
        if self._resume_word_count(resume_json) > RESUME_HARD_WORD_CAP:
            warnings.append("hard_length_risk")
        if self._has_run_together_school_date(resume_json):
            warnings.append("run_together_school_date")
        if self._selected_coursework(resume_json) and not resume_json.get("education"):
            warnings.append("coursework_without_education")
        if resume_json.get("skills") and not self._group_skills(
            resume_json.get("skills", []),
            separate_programming_data=self._is_software_data_heavy_resume(resume_json),
            resume_json=resume_json,
        ):
            warnings.append("skills_not_grouped")
        if "?" in self._render_resume_text(resume_json):
            warnings.append("replacement_character_detected")
        if self._education_has_deans_list(resume_json) and not self._deans_list_preserved(resume_json):
            warnings.append("deans_list_missing")
        utilization = self._resume_page_utilization(resume_json)
        if utilization < RESUME_PAGE_UTILIZATION_TARGET and not resume_json.get("rendering_allocation", {}).get("one_page_enforced"):
            warnings.append("low_page_utilization")
        estimated_page_count = self._resume_estimated_page_count(resume_json)
        if estimated_page_count > 1:
            warnings.append("page_count_exceeded")
        missing_major = self._missing_major_sections(resume_json, sections)
        warnings.extend(f"missing_major_section:{section}" for section in missing_major)
        if not self._expansion_order_is_reasonable(resume_json):
            warnings.append("expansion_order_risk")
        return {
            "ok": not warnings,
            "warnings": warnings,
            "sections": sections,
            "page_utilization_estimate": utilization,
            "estimated_page_count": estimated_page_count,
            "major_sections_present": [section for section in self._required_major_sections(resume_json) if section in sections],
            "expansion_order": resume_json.get("rendering_allocation", {}).get("expansion_order", []),
        }

    def _planned_resume_sections(self, resume_json: dict) -> list[str]:
        sections = []
        if resume_json.get("professional_summary"):
            sections.append("Professional Summary")
        if resume_json.get("education"):
            sections.append("Education")
        if self._selected_coursework(resume_json):
            sections.append("Relevant Coursework")
        if resume_json.get("projects"):
            sections.append("Engineering Experience")
        if resume_json.get("experience"):
            sections.append("Work Experience")
        if self._group_skills(
            resume_json.get("skills", []),
            separate_programming_data=self._is_software_data_heavy_resume(resume_json),
            resume_json=resume_json,
        ):
            sections.append("Technical Skills")
        if resume_json.get("certifications"):
            sections.append("Professional Development")
        return sections

    @staticmethod
    def _education_has_deans_list(resume_json: dict) -> bool:
        return any("dean" in " ".join(map(str, edu.get("honors", []))).lower() for edu in resume_json.get("education", []))

    @staticmethod
    def _deans_list_preserved(resume_json: dict) -> bool:
        text = json.dumps(resume_json).lower()
        return "dean" in text

    def _resume_page_utilization(self, resume_json: dict) -> float:
        return round(min(self._resume_word_count(resume_json) / RESUME_SOFT_WORD_CAP, 1.25), 2)

    def _resume_estimated_page_count(self, resume_json: dict) -> int:
        if self._resume_word_count(resume_json) > RESUME_ONE_PAGE_WORD_CAP:
            return 2
        if len(self._selected_coursework(resume_json)) > RESUME_ONE_PAGE_COURSEWORK_CAP:
            return 2
        if len(self._clean_string_list(resume_json.get("skills", []))) > RESUME_ONE_PAGE_SKILL_CAP:
            return 2
        if len(self._clean_string_list(resume_json.get("certifications", []))) > 1:
            return 2
        if len(self._clean_text(resume_json.get("professional_summary", "")).split()) > RESUME_ONE_PAGE_SUMMARY_WORD_CAP:
            return 2
        for project in resume_json.get("projects", []):
            bullets = project.get("bullets", [])
            if self._is_capstone_project(project) and len(bullets) > 4:
                return 2
            if self._is_job_search_assistant_project(project) and len(bullets) > 1:
                return 2
            if not self._is_capstone_project(project) and not self._is_job_search_assistant_project(project) and len(bullets) > 1:
                return 2
        for entry in resume_json.get("experience", []):
            if len(entry.get("bullets", [])) > 2:
                return 2
        return 1

    @staticmethod
    def _required_major_sections(resume_json: dict) -> list[str]:
        required = ["Education", "Engineering Experience", "Technical Skills"]
        if resume_json.get("experience"):
            required.append("Work Experience")
        if resume_json.get("certifications"):
            required.append("Professional Development")
        return required

    def _missing_major_sections(self, resume_json: dict, sections: list[str]) -> list[str]:
        return [section for section in self._required_major_sections(resume_json) if section not in sections]

    @staticmethod
    def _expansion_order_is_reasonable(resume_json: dict) -> bool:
        order = resume_json.get("rendering_allocation", {}).get("expansion_order", [])
        if not order:
            return True
        priority = [
            "capstone_bullets",
            "job_search_assistant_bullets",
            "academic_project_bullets",
            "work_experience_bullets",
            "coursework_items",
            "skills_items",
            "professional_development_items",
            "summary_detail",
        ]
        ranked = [priority.index(item) for item in order if item in priority]
        return ranked == sorted(ranked)

    @staticmethod
    def _has_run_together_school_date(resume_json: dict) -> bool:
        for edu in resume_json.get("education", []):
            institution = str(edu.get("institution", "")).strip()
            graduation = str(edu.get("graduation", "")).strip()
            if institution and graduation and f"{institution}{graduation}" in json.dumps(edu):
                return True
        return False

    def _selected_coursework(self, resume_json: dict) -> list[str]:
        coursework: list[str] = []
        for edu in resume_json.get("education", []):
            coursework.extend(self._clean_string_list(edu.get("relevant_coursework", [])))
        limit = MAX_RESUME_COURSEWORK_EXPANDED if resume_json.get("rendering_allocation", {}).get("expanded_coursework") else MAX_RESUME_COURSEWORK
        return self._arrange_coursework_for_table(self._compact_coursework_list(coursework)[:limit])

    @classmethod
    def _arrange_coursework_for_table(cls, coursework: list[str]) -> list[str]:
        compacted = cls._compact_coursework_list(coursework)
        return sorted(compacted, key=lambda value: (len(value) > 28, len(value), value.lower()))

    @classmethod
    def _compact_coursework_list(cls, coursework: list[str]) -> list[str]:
        compacted = []
        seen = set()
        for course in coursework:
            label = cls._compact_coursework_label(course)
            key = label.lower()
            if key not in seen:
                compacted.append(label)
                seen.add(key)
        return compacted

    @staticmethod
    def _compact_coursework_label(value: str) -> str:
        text = DocumentGenerator._clean_text(value)
        normalized = text.lower().replace("&", "and")
        replacements = {
            "field practices in civil engineering technology": "Field Practices in Civil Engineering",
            "field practices in civil engineering": "Field Practices in Civil Engineering",
            "soils and foundations": "Soils & Foundations",
            "soil mechanics and foundations": "Soils & Foundations",
            "soil mechanics & foundations": "Soils & Foundations",
        }
        return replacements.get(normalized, text)

    def _profile_education_entry(self, institution: str) -> dict | None:
        if not self._profile or not institution:
            return None
        normalized = str(institution).strip().lower()
        for entry in self._profile.get("education", []) or []:
            if str(entry.get("institution", "")).strip().lower() == normalized:
                return entry
        for key, entry in (self._profile.get("education_detail", {}) or {}).items():
            if not isinstance(entry, dict):
                continue
            key_text = str(key).replace("_", " ").lower()
            entry_name = str(entry.get("institution", "")).strip().lower()
            if normalized == entry_name or all(part in key_text for part in normalized.split()[:2]):
                return entry
        return None

    def _education_accreditation(self, edu: dict) -> str:
        candidates = []
        institution = str(edu.get("institution", "")).strip().lower()
        if "farmingdale" in institution:
            candidates.append("ABET Accredited")
        if edu.get("accreditation"):
            candidates.append(edu.get("accreditation"))
        profile_edu = self._profile_education_entry(edu.get("institution", ""))
        if profile_edu and profile_edu.get("accreditation"):
            candidates.append(profile_edu.get("accreditation"))
        for candidate in candidates:
            if isinstance(candidate, dict):
                text = " ".join(str(value) for value in candidate.values() if value)
            else:
                text = str(candidate)
            if "abet" in text.lower():
                return "ABET Accredited"
        return ""

    def _education_location(self, edu: dict) -> str:
        if edu.get("location"):
            return str(edu["location"]).strip()
        city = edu.get("city") or edu.get("location_city")
        state = edu.get("state") or edu.get("location_state")
        location = ", ".join([str(x).strip() for x in [city, state] if x])
        if location:
            return location
        profile_edu = self._profile_education_entry(edu.get("institution", ""))
        if profile_edu:
            if profile_edu.get("location"):
                return str(profile_edu["location"]).strip()
            city = profile_edu.get("city") or profile_edu.get("location_city")
            state = profile_edu.get("state") or profile_edu.get("location_state")
            location = ", ".join([str(x).strip() for x in [city, state] if x])
            if location:
                return location
        return SCHOOL_LOCATION_FALLBACKS.get(str(edu.get("institution", "")).strip().lower(), "")

    def _degree_line(self, edu: dict) -> str:
        degree = str(edu.get("degree", "")).strip()
        major = str(edu.get("major", "")).strip()
        accreditation = self._education_accreditation(edu)
        if degree and major:
            line = f"{degree} in {major}"
        else:
            line = degree or major
        if accreditation and "abet" not in line.lower():
            line = f"{line} ({accreditation})" if line else accreditation
        line = self._dedupe_abet_references(line)
        return line

    @staticmethod
    def _education_graduation_date(value: str) -> str:
        text = DocumentGenerator._clean_text(value)
        match = re.search(
            r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b",
            text,
            flags=re.IGNORECASE,
        )
        if not match:
            return text
        return " ".join(part.capitalize() if idx == 0 else part for idx, part in enumerate(match.group(0).split()))

    @staticmethod
    def _education_honors(edu: dict) -> list[str]:
        return [honor for honor in DocumentGenerator._clean_string_list(edu.get("honors", [])) if "abet" not in honor.lower()]

    @staticmethod
    def _dedupe_abet_references(value: str) -> str:
        text = re.sub(r"\s+", " ", value).strip()
        if "abet" not in text.lower():
            return text
        text = re.sub(r"\s*\([^)]*abet[^)]*\)", "", text, flags=re.IGNORECASE).strip()
        text = re.sub(r"\s*[-\u2013\u2014,;:|]*\s*ABET[-\s]*Accredited\b", "", text, flags=re.IGNORECASE).strip()
        return f"{text} (ABET Accredited)" if text else "ABET Accredited"

    def _group_skills(
        self,
        skills: list[str],
        separate_programming_data: bool = False,
        resume_json: dict | None = None,
    ) -> dict[str, list[str]]:
        groups: dict[str, list[str]] = {
            "Software": [],
            "Engineering": [],
            "Codes/Standards": [],
        }
        if separate_programming_data:
            groups["Programming/Data"] = []
        for skill in self._clean_string_list(skills):
            normalized = skill.lower()
            if any(term in normalized for term in ["python", "sqlite", "openai", "api", "llm", "data", "cli", "google sheets", "google drive", "automation"]):
                target_group = "Programming/Data" if separate_programming_data else "Software"
                groups[target_group].append(skill)
            elif any(term in normalized for term in ["asce", "aisc", "aci", "astm", "osha", "code", "standard"]):
                groups["Codes/Standards"].append(self._canonical_standard_skill(skill))
            elif any(term in normalized for term in ["autocad", "revit", "ram", "excel", "matlab", "project", "inventor", "civil 3d", "bluebeam"]):
                groups["Software"].append(skill)
            else:
                groups["Engineering"].append(skill)
        for standard in self._supported_standard_skills(resume_json):
            if not any(self._standard_key(standard) == self._standard_key(existing) for existing in groups["Codes/Standards"]):
                groups["Codes/Standards"].append(standard)
        groups["Codes/Standards"] = self._dedupe_standards(groups["Codes/Standards"])
        return {label: values for label, values in groups.items() if values}

    def _supported_standard_skills(self, resume_json: dict | None = None) -> list[str]:
        if not self._profile:
            return []
        evidence = " ".join(
            self._clean_string_list((self._profile.get("codes_and_standards", {}) or {}).get("verified_use_or_awareness", []))
        )
        if not evidence:
            return []
        resume_text = json.dumps(resume_json or {}).lower()
        structurally_relevant = any(
            term in resume_text
            for term in ("structural", "steel", "concrete", "foundation", "framing", "ram", "civil")
        )
        if not structurally_relevant:
            return []
        standards: list[str] = []
        supported = evidence.lower()
        for label, signal in [
            ("ASCE 7", "asce"),
            ("AISC Steel Construction Manual", "aisc"),
            ("ACI 318", "aci"),
            ("ASTM D854", "astm d854"),
        ]:
            if signal in supported:
                standards.append(label)
        return standards

    @classmethod
    def _dedupe_standards(cls, standards: list[str]) -> list[str]:
        out: list[str] = []
        seen = set()
        for standard in standards:
            canonical = cls._canonical_standard_skill(standard)
            key = cls._standard_key(canonical)
            if key and key not in seen:
                out.append(canonical)
                seen.add(key)
        return out

    @staticmethod
    def _canonical_standard_skill(value: str) -> str:
        text = DocumentGenerator._clean_text(value)
        normalized = text.lower()
        if "asce 7" in normalized or normalized.startswith("asce"):
            return "ASCE 7"
        if "aisc" in normalized:
            return "AISC Steel Construction Manual"
        if "aci 318" in normalized or normalized.startswith("aci"):
            return "ACI 318"
        if "astm d854" in normalized:
            return "ASTM D854"
        return text

    @staticmethod
    def _standard_key(value: str) -> str:
        normalized = DocumentGenerator._clean_text(value).lower()
        if "asce 7" in normalized or normalized.startswith("asce"):
            return "asce 7"
        if "aisc" in normalized:
            return "aisc"
        if "aci 318" in normalized or normalized.startswith("aci"):
            return "aci 318"
        if "astm d854" in normalized:
            return "astm d854"
        return normalized

    def _project_organization(self, project: dict) -> str:
        if self._is_job_search_assistant_project(project):
            return "Personal Project"
        if self._is_bridge_construction_project(project):
            return "Farmingdale State College"
        for key in ("organization", "program", "team", "context"):
            value = self._clean_text(project.get(key, ""))
            if value:
                return value
        project_type = self._clean_text(project.get("type", ""))
        if project_type:
            if "personal" in project_type.lower():
                return "Personal Project"
            return project_type.replace("_", " ").title()
        if self._is_job_search_assistant_project(project):
            return "Personal Project"
        if self._is_academic_project(project):
            return "Farmingdale State College"
        return ""

    def _project_role(self, project: dict) -> str:
        if self._is_job_search_assistant_project(project):
            return "Project Architect / Lead Developer"
        if self._is_bridge_construction_project(project):
            return "Construction Planning Project"
        role = self._clean_text(project.get("role", ""))
        if role and role.lower() not in {"academic project", "class project", "student project", "personal project"}:
            return role
        if self._is_academic_project(project):
            return "Student Project Contributor"
        if self._is_personal_project(project):
            return "Personal Project Contributor"
        return role

    def _project_date(self, project: dict) -> str:
        date = self._clean_text(project.get("date", ""))
        if date and self._looks_like_project_date(date):
            return date
        text = json.dumps(project).lower()
        if self._is_personal_project(project):
            return "2024"
        if any(term in text for term in ("volunteer", "community", "cleanup", "asce student chapter")):
            return "2024"
        return "2025"

    @staticmethod
    def _is_personal_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return "personal" in text or "job search assistant" in text or "software automation" in text

    @staticmethod
    def _is_academic_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return any(term in text for term in ("academic", "student", "capstone", "class_project", "course", "farmingdale", "bridge replacement"))

    @staticmethod
    def _is_bridge_construction_project(project: dict) -> bool:
        text = json.dumps(project).lower()
        return "bridge" in text and any(term in text for term in ("construction", "management", "planning", "plan"))

    @staticmethod
    def _looks_like_project_date(value: str) -> bool:
        text = value.strip().lower()
        if not text:
            return False
        if any(label in text for label in ("academic project", "personal project", "student project", "class project")):
            return False
        return any(char.isdigit() for char in text) or any(term in text for term in ("spring", "summer", "fall", "winter"))

    @staticmethod
    def _remove_table_borders(table) -> None:
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        borders = tbl_pr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = f"w:{edge}"
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            element.set(qn("w:val"), "nil")

    @staticmethod
    def _add_right_tab(paragraph) -> None:
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(RESUME_RIGHT_TAB_INCHES),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.SPACES,
        )

    @staticmethod
    def _set_spacing(paragraph, before: float | None = None, after: float | None = None, line: float | None = None) -> None:
        paragraph.paragraph_format.space_before = Pt(before or 0)
        paragraph.paragraph_format.space_after = Pt(after or 0)
        if line is not None:
            paragraph.paragraph_format.line_spacing = line

    @staticmethod
    def _format_run(run, size: int = 10, bold: bool | None = None, italic: bool | None = None) -> None:
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic

    def _format_paragraph_runs(self, paragraph, size: int = 10) -> None:
        for run in paragraph.runs:
            if run.font.size is None:
                run.font.size = Pt(size)
            if run.font.name is None:
                run.font.name = "Calibri"

    @staticmethod
    def _add_bottom_border(paragraph) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        p_bdr.append(bottom)

    def _add_body_paragraph(self, doc, text: str):
        p = doc.add_paragraph(text)
        self._set_spacing(p)
        self._format_paragraph_runs(p, size=10)
        return p

    def _add_cover_paragraph(self, doc, text: str, before: float = 0, after: float = COVER_BODY_AFTER_PT):
        p = doc.add_paragraph(text)
        self._set_spacing(p, before=before, after=after, line=COVER_LINE_SPACING)
        self._format_paragraph_runs(p, size=11)
        return p

    def _add_bullet(self, doc, text: str):
        p = doc.add_paragraph(style="Resume Bullet")
        self._set_spacing(p)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(self._clean_bullet_text(text))
        self._format_run(run, size=10)
        return p

    @staticmethod
    def _clean_bullet_text(text: str) -> str:
        cleaned = DocumentGenerator._clean_text(text)
        return re.sub(r"\s*[|;:,]+\s*$", "", cleaned).strip()

    def _clean_cover_closing(self, value: str) -> str:
        text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        self._reject_template_placeholders(text, "cover_letter_closing")
        if not text:
            text = "Sincerely,\nJames Morseman"
        lines = [DocumentGenerator._clean_text(line) for line in text.split("\n") if DocumentGenerator._clean_text(line)]
        if len(lines) == 1:
            match = re.match(r"^(Sincerely|Regards|Best regards|Respectfully),?\s+(.+)$", lines[0], flags=re.IGNORECASE)
            if match:
                closing = match.group(1)
                name = match.group(2)
                lines = [f"{closing[0].upper()}{closing[1:]},", name]
        signature_name = self._cover_signature_name()
        if lines:
            if len(lines) == 1:
                lines.append(signature_name)
            else:
                lines[-1] = signature_name
        return "\n".join(lines)

    def _cover_signature_name(self) -> str:
        identity = (self._profile or {}).get("identity", {})
        first = self._clean_text(identity.get("first_name", "")) or "James"
        last = self._clean_text(identity.get("last_name", "")) or "Morseman"
        return f"{first} {last}".strip()

    @staticmethod
    def _contains_template_placeholder(value: str) -> bool:
        return bool(TEMPLATE_PLACEHOLDER_RE.search(value or ""))

    @staticmethod
    def _contains_unresolved_template_marker(value: str) -> bool:
        text = value or ""
        return bool(
            TEMPLATE_PLACEHOLDER_RE.search(text)
            or UNRESOLVED_TEMPLATE_MARKER_RE.search(text)
        )

    @staticmethod
    def _reject_template_placeholders(value: str, context: str) -> None:
        if DocumentGenerator._contains_unresolved_template_marker(value):
            raise ValueError(f"Template placeholder or marker leaked into {context}.")
